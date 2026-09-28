"""Design helpers for the v7 deck, built on the FKS Food & Agri template with the v6 palette."""
import json, os, math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, 'assets')

# ---- v6 palette (measured from the v6 PDF) ----
DARK = '4D4D4F'      # titles, body, dark pictograms
MID = '747678'       # secondary text, subtitles
GREY = 'A7A9AC'      # light grey pictograms / conventional
GREY2 = '6D6E70'     # darker grey pictograms
LINE = 'D9DADC'      # rules, map land
PANEL = 'F1F1F2'     # light grey panels
YELLOW = 'E6B222'    # FKS accent
YELLOW2 = 'F7DD59'   # light yellow (pictogram highlights)
CREAM = 'FDF8E7'     # FKS 'after' panels
CREAM2 = 'FBF0CC'    # yellow tint highlight
GOLD = '8A6A0C'      # dark gold text
GOLD2 = 'A8820F'
WHITE = 'FFFFFF'
FONT = 'Calibri'

W, H = 10.0, 5.625
ML, MR = 0.56, 9.44   # content margins

def rgb(hexstr):
    return RGBColor.from_string(hexstr)

def I(v):
    return Inches(v)

# ---------------------------------------------------------------- text
def text(slide, s, x, y, w, h, size=10, bold=False, color=DARK, align='left', anchor='top',
         italic=False, spacing=None, line_spacing=None, font=FONT, wrap=True, name=None):
    """Add a text box. `s` may be a string or a list of runs: [(text, {bold, color, size, italic}), ...].
    Use '\n' in strings for new paragraphs."""
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}[anchor]
    paras = s.split('\n') if isinstance(s, str) else [s]
    for pi, ptxt in enumerate(paras):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
        if line_spacing:
            p.line_spacing = line_spacing
        runs = [(ptxt, {})] if isinstance(ptxt, str) else ptxt
        for rtxt, ro in runs:
            r = p.add_run()
            r.text = rtxt
            f = r.font
            f.name = font
            f.size = Pt(ro.get('size', size))
            f.bold = ro.get('bold', bold)
            f.italic = ro.get('italic', italic)
            f.color.rgb = rgb(ro.get('color', color))
            sp = ro.get('spacing', spacing)
            if sp is not None:
                r._r.get_or_add_rPr().set('spc', str(int(sp * 100)))
    return tb

def label(slide, s, x, y, w, color=MID, size=9):
    """Section label in the v6 style: small caps, letter-spaced, bold."""
    return text(slide, s.upper(), x, y, w, 0.22, size=size, bold=True, color=color, spacing=1.5)

# ---------------------------------------------------------------- shapes
def _style(shape, fill=None, line=None, lw=0.75, dash=None):
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line); shape.line.width = Pt(lw)
        if dash:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shape.shadow.inherit = False
    return shape

def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=0.75, dash=None, rot=0, flipH=False, flipV=False, name=None, adj=None):
    shp = slide.shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    _style(shp, fill, line, lw, dash)
    if rot:
        shp.rotation = rot
    if flipH or flipV:
        xfrm = shp._element.spPr.xfrm
        if flipH: xfrm.set('flipH', '1')
        if flipV: xfrm.set('flipV', '1')
    if adj is not None:
        for i, v in enumerate(adj):
            shp.adjustments[i] = v
    if name:
        shp.name = name
    # strip default text placeholder styling: no text
    return shp

def rect(slide, x, y, w, h, fill=PANEL, line=None, lw=0.75, name=None):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, lw, name=name)

def rrect(slide, x, y, w, h, fill=PANEL, line=None, lw=0.75, radius=0.1, name=None):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, lw, name=name)
    s.adjustments[0] = min(0.5, radius / min(w, h))
    return s

def oval(slide, x, y, w, h, fill=YELLOW, line=None, lw=0.75, name=None):
    return shape(slide, MSO_SHAPE.OVAL, x, y, w, h, fill, line, lw, name=name)

def line(slide, x1, y1, x2, y2, color=LINE, lw=0.75, dash=None, arrow=False, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn('a:tailEnd')); tail.set('type', 'triangle'); tail.set('w', 'med'); tail.set('len', 'med')
    if name:
        c.name = name
    return c

def chevron(slide, x, y, w, h, fill=YELLOW):
    return shape(slide, MSO_SHAPE.CHEVRON, x, y, w, h, fill, None)

def tri_marker(slide, x, y, size=0.12, fill=YELLOW):
    """The small yellow 'play' triangle the template uses as a bullet."""
    return shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, x, y, size * 0.85, size, fill, None, rot=90)

def picture(slide, path, x, y, w, h, crop_cover=True, name=None):
    """Insert an image; crop to cover the box (keeps aspect) using PIL first."""
    from PIL import Image
    src = path
    if crop_cover:
        im = Image.open(path)
        iw, ih = im.size
        target = w / h
        if iw / ih > target:
            nw = int(ih * target); left = (iw - nw) // 2; im = im.crop((left, 0, left + nw, ih))
        else:
            nh = int(iw / target); top = (ih - nh) // 2; im = im.crop((0, top, iw, top + nh))
        src = os.path.join(ASSETS, '_crop_%s_%dx%d.png' % (os.path.basename(path).split('.')[0], int(w * 100), int(h * 100)))
        im.save(src)
    pic = slide.shapes.add_picture(src, I(x), I(y), I(w), I(h))
    if name:
        pic.name = name
    return pic

# ---------------------------------------------------------------- map
POLYS = json.load(open(os.path.join(ASSETS, 'map_polys.json')))
NEIGHBOURS = ['Malaysia', 'Singapore', 'Brunei', 'East Timor', 'Papua New Guinea', 'Philippines', 'Australia', 'Thailand', 'Vietnam', 'Cambodia']
ID_BBOX = (94.5, -11.5, 141.5, 6.5)
WEST_BBOX = (94.5, -9.5, 120.5, 6.5)

def map_height(w, bbox):
    lon0, lat0, lon1, lat1 = bbox
    mean = (lat0 + lat1) / 2 * math.pi / 180
    return w * (lat1 - lat0) / ((lon1 - lon0) * math.cos(mean))

def projector(x, y, w, h, bbox):
    lon0, lat0, lon1, lat1 = bbox
    return lambda lon, lat: (x + (lon - lon0) / (lon1 - lon0) * w, y + (lat1 - lat) / (lat1 - lat0) * h)

def map_layer(slide, which, x, y, w, h, bbox, fill=LINE, line_color=WHITE, lw=0.5, name=None):
    """Editable freeform coastline for Indonesia or its neighbours, clipped to the box."""
    lon0, lat0, lon1, lat1 = bbox
    countries = [which] if which != 'neighbours' else NEIGHBOURS
    E = 914400
    ff = None
    for c in countries:
        for ring in POLYS.get(c, []):
            pts = []
            for lon, lat in ring:
                X = (lon - lon0) / (lon1 - lon0) * w; Y = (lat1 - lat) / (lat1 - lat0) * h
                pts.append((X, Y))
            if all(p[0] < 0 or p[0] > w or p[1] < 0 or p[1] > h for p in pts):
                continue
            pts = [(min(max(X, 0), w), min(max(Y, 0), h)) for X, Y in pts]
            dd = []
            for p in pts:
                q = (int(round(p[0] * E)), int(round(p[1] * E)))
                if not dd or abs(dd[-1][0] - q[0]) > 1 or abs(dd[-1][1] - q[1]) > 1:
                    dd.append(q)
            if len(dd) < 3:
                continue
            if ff is None:
                ff = slide.shapes.build_freeform(dd[0][0], dd[0][1], scale=1.0)
            else:
                ff.move_to(dd[0][0], dd[0][1])
            ff.add_line_segments(dd[1:], close=True)
    if ff is None:
        return None
    shp = ff.convert_to_shape(I(x), I(y))
    _style(shp, fill, line_color, lw)
    shp.name = name or ('Map - ' + ('Indonesia' if which == 'Indonesia' else 'Neighbouring countries'))
    return shp

def draw_map(slide, x, y, w, bbox, land=LINE, neighbours='E8E9EA'):
    h = map_height(w, bbox)
    map_layer(slide, 'neighbours', x, y, w, h, bbox, fill=neighbours, line_color=WHITE, lw=0.5)
    map_layer(slide, 'Indonesia', x, y, w, h, bbox, fill=land, line_color=WHITE, lw=0.5)
    return projector(x, y, w, h, bbox), h

# ---------------------------------------------------------------- pictograms (v6 shape language)
def ship(slide, x, y, w, h, color=GREY, hatches=4, aft_left=False, hatch_color=None):
    shape(slide, MSO_SHAPE.TRAPEZOID, x, y + h * 0.58, w, h * 0.42, color, None, flipV=True)
    hx = x + w * 0.1 if aft_left else x + w * 0.74
    shape(slide, MSO_SHAPE.RECTANGLE, hx, y + h * 0.28, w * 0.16, h * 0.32, color, None)
    shape(slide, MSO_SHAPE.RECTANGLE, hx + w * 0.03, y + h * 0.1, w * 0.1, h * 0.2, color, None)
    hw = (w * 0.62) / hatches
    h0 = x + w * 0.3 if aft_left else x + w * 0.08
    for i in range(hatches):
        shape(slide, MSO_SHAPE.RECTANGLE, h0 + i * hw + hw * 0.1, y + h * 0.44, hw * 0.8, h * 0.14, hatch_color or color, None)

def crane(slide, x, y, w, h, color=GREY):
    shape(slide, MSO_SHAPE.RECTANGLE, x + w * 0.12, y, w * 0.07, h, color, None)
    shape(slide, MSO_SHAPE.RECTANGLE, x + w * 0.12, y, w * 0.8, h * 0.07, color, None)
    line(slide, x + w * 0.8, y + h * 0.07, x + w * 0.8, y + h * 0.52, color, 1.0)
    shape(slide, MSO_SHAPE.FLOWCHART_MERGE, x + w * 0.66, y + h * 0.52, w * 0.28, h * 0.3, color, None)

def hopper(slide, x, y, w, h, color=GREY):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h * 0.32, color, None)
    shape(slide, MSO_SHAPE.FLOWCHART_MERGE, x + w * 0.08, y + h * 0.3, w * 0.84, h * 0.7, color, None)
    shape(slide, MSO_SHAPE.RECTANGLE, x + w * 0.1, y + h * 0.3, w * 0.05, h * 0.7, color, None)
    shape(slide, MSO_SHAPE.RECTANGLE, x + w * 0.85, y + h * 0.3, w * 0.05, h * 0.7, color, None)

def truck(slide, x, y, w, h, color=GREY, cab_color=None):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y + h * 0.1, w * 0.66, h * 0.55, color, None)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + w * 0.68, y + h * 0.3, w * 0.3, h * 0.35, cab_color or color, None)
    d = h * 0.34
    for fx in (0.1, 0.36, 0.74):
        shape(slide, MSO_SHAPE.OVAL, x + w * fx, y + h * 0.66, d, d, color, None)

def mill(slide, x, y, w, h, color=DARK):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y + h * 0.42, w, h * 0.58, color, None)
    n = 3; tw = w / n
    for i in range(n):
        shape(slide, MSO_SHAPE.RIGHT_TRIANGLE, x + i * tw, y + h * 0.16, tw, h * 0.27, color, None, flipH=True)
    shape(slide, MSO_SHAPE.RECTANGLE, x + w * 0.8, y, w * 0.1, h * 0.3, color, None)

def silos(slide, x, y, w, h, n=3, color=DARK):
    gap = w * 0.06; sw = (w - gap * (n - 1)) / n
    for i in range(n):
        shape(slide, MSO_SHAPE.CAN, x + i * (sw + gap), y, sw, h, color, None)

def warehouse(slide, x, y, w, h, color=YELLOW, door=True):
    shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, x - w * 0.04, y, w * 1.08, h * 0.32, color, None)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y + h * 0.3, w, h * 0.7, color, None)
    if door:
        shape(slide, MSO_SHAPE.RECTANGLE, x + w * 0.42, y + h * 0.62, w * 0.16, h * 0.38, WHITE, None)

def conveyor(slide, x, y, w, color=YELLOW, thickness=0.07, legs=3, leg_h=0.45):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, thickness, color, None)
    for i in range(legs):
        lx = x + (w - 0.03) * (i / max(1, legs - 1)) + 0.015
        shape(slide, MSO_SHAPE.RECTANGLE, lx - 0.012, y + thickness, 0.024, leg_h, color, None)

def unloader(slide, x, y, w, h, color=YELLOW):
    shape(slide, MSO_SHAPE.RECTANGLE, x + w * 0.55, y + h * 0.15, w * 0.07, h * 0.85, color, None)
    shape(slide, MSO_SHAPE.RECTANGLE, x + w * 0.85, y + h * 0.15, w * 0.07, h * 0.85, color, None)
    shape(slide, MSO_SHAPE.RECTANGLE, x + w * 0.5, y + h * 0.12, w * 0.5, h * 0.1, color, None)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y + h * 0.02, w * 0.62, h * 0.08, color, None)
    shape(slide, MSO_SHAPE.RECTANGLE, x + w * 0.08, y + h * 0.08, w * 0.06, h * 0.5, color, None)
    shape(slide, MSO_SHAPE.RECTANGLE, x + w * 0.62, y + h * 0.02, w * 0.05, h * 0.2, color, None)

# ---------------------------------------------------------------- notes
def notes(slide, d):
    parts = [
        'SLIDE OBJECTIVE\n' + d['objective'].strip(),
        'TALK TRACK\n' + d['talk'].strip(),
        'SOURCE / DATA BASIS\n' + d['sources'].strip(),
        'CALCULATION BASIS\n' + d['calc'].strip(),
        'INTERNAL DATA STILL TO VALIDATE\n' + d['validate'].strip(),
    ]
    slide.notes_slide.notes_text_frame.text = '\n\n'.join(parts)
