"""Builds the three sheets of a project (Input / Calc / Output) from a SPEC dict.

SPEC schema (see specs/p02_cigconveyor.py for a worked example):
  id, code, name, short, entity, category, description, v29 (dict of references)
  defaults: in_valuation (1/0), case (1 = both cases, 2 = With-BUMN case only), econ_interest (0-1)
  commissioning: dict(year, source, cls, note) or None for existing businesses (operating from 2024)
  capex: dict(total_usd, source, pct {year: share}, pct_source, building_share, building_source,
              life_b, life_e, lives_source, hist_confirmed (0/1), note) or None (no project capex)
  inputs: list of rows  dict(key, label, unit, kind='annual'|'scalar', values={year: v} | value, source, cls,
                             note, fmt, section)
  calc:   list of rows  dict(key, label, unit, f, f0, fmt, section, total, bold, note, src)
  hooks:  dict(revenue, cogs, opex, da_other, other_income)  -> calc keys (cogs/opex/da positive cost amounts)
  depends: list of project ids whose calc rows are referenced through {x:PID:key}
  existing: True for existing businesses (no entry cost; opening WC allocated)
"""
from openpyxl.styles import Alignment
from . import styles as S
from .engine import (Model, YEARS, FORECAST_YEARS, YEAR_COL, FIRST_COL, LAST_COL, FCOL, LCOL, L, yc,
                     DATA_ROW0, SCALAR_COL, SRC_COL, NOTE_COL, LABEL_COL, UNIT_COL, TYPE_COL, ROW_YEAR,
                     CLASS_LABEL, CLASS_FILL, GLOBAL_TITLE)

TAB_IN, TAB_CALC, TAB_OUT = "FFD966", "BFBFBF", "A9D08E"


def sheet_titles(spec):
    if spec.get("titles"):
        return tuple(spec["titles"])
    base = f"{spec['id']} {spec['code']}"
    return (f"{base} Input", f"{base} Calc", f"{base} Output")


# ----------------------------------------------------------------------------------------------- planning
class ProjectPlan:
    """Assigns row numbers for a project's three sheets and registers them (phase 1)."""

    def __init__(self, model: Model, spec):
        self.m = model
        self.spec = spec
        self.pid = spec["id"]
        self.kin, self.kcalc, self.kout = f"{self.pid}:in", f"{self.pid}:calc", f"{self.pid}:out"
        self.in_rows, self.calc_rows, self.out_rows = [], [], []   # (type, payload, row)
        t_in, t_calc, t_out = sheet_titles(spec)
        model.add_sheet(self.kin, t_in, TAB_IN)
        model.add_sheet(self.kcalc, t_calc, TAB_CALC)
        model.add_sheet(self.kout, t_out, TAB_OUT)
        model.projects[self.pid] = spec
        self._plan_input()
        self._plan_calc()
        self._plan_output()

    # ---- helpers
    def _add(self, rows, typ, payload, sheetkey, row, key=None, kind="annual", fmt=None):
        rows.append((typ, payload, row))
        if key:
            self.m.register(sheetkey, key, row, kind=kind, fmt=fmt, label=payload.get("label", "") if isinstance(payload, dict) else "")
        return row + 1

    # ---- input sheet
    def _plan_input(self):
        sp = self.spec
        r = DATA_ROW0
        r = self._add(self.in_rows, "section", {"text": "1. Project register (identification)"}, self.kin, r)
        for lab, val in [("Project ID", sp["id"]), ("Project", sp["name"]), ("Legal entity / operator", sp.get("entity", "")),
                         ("Category", sp.get("category", "")), ("Scope", sp.get("description", ""))]:
            r = self._add(self.in_rows, "text", {"label": lab, "value": val}, self.kin, r)
        v29 = sp.get("v29", {})
        for lab, k in [("v29 calculation block", "calc"), ("v29 input block", "input"),
                       ("v29 standalone P&L block", "standalone"), ("v29 inclusion flag", "flag")]:
            if v29.get(k):
                r = self._add(self.in_rows, "text", {"label": lab, "value": v29[k]}, self.kin, r)
        r += 1
        r = self._add(self.in_rows, "section", {"text": "2. Controls (company-level settings — edit on Global Inputs; shown here for reference)"}, self.kin, r)
        r = self._add(self.in_rows, "link", {"label": "Included in valuation (1 = yes, 0 = no)", "unit": "0/1", "f": "={g:incl." + self.pid + "}", "fmt": "flag01",
                                             "src": "Global Inputs — project register"}, self.kin, r, key="incl", kind="scalar", fmt="flag01")
        r = self._add(self.in_rows, "link", {"label": "Case membership (1 = both cases, 2 = With-BUMN case only)", "unit": "1/2", "f": "={g:case." + self.pid + "}", "fmt": "flag01",
                                             "src": "Global Inputs — project register"}, self.kin, r, key="case", kind="scalar", fmt="flag01")
        r = self._add(self.in_rows, "link", {"label": "FSL economic interest", "unit": "%", "f": "={g:eint." + self.pid + "}", "fmt": "pct",
                                             "src": "Global Inputs — project register"}, self.kin, r, key="eint", kind="scalar", fmt="pct")
        com = sp.get("commissioning")
        if com:
            r = self._add(self.in_rows, "input", {"key": "commissioning", "label": "Commissioning / first operating (revenue) year", "unit": "year",
                                                  "kind": "scalar", "value": com["year"], "source": com.get("source", ""), "cls": com.get("cls", "v29"),
                                                  "note": com.get("note", "Operations start in this year; capex timing is set separately below."), "fmt": "year"},
                          self.kin, r, key="commissioning", kind="scalar", fmt="year")
        r += 1
        # spec inputs grouped by section
        cur = None
        n = 3
        for it in sp.get("inputs", []):
            sec = it.get("section", "Operating assumptions")
            if sec != cur:
                cur = sec
                r = self._add(self.in_rows, "section", {"text": f"{n}. {sec}"}, self.kin, r)
                n += 1
            kind = it.get("kind", "annual")
            fmt = it.get("fmt") or S.UNIT_FMT.get(it.get("unit", ""), "num2")
            r = self._add(self.in_rows, "input", it, self.kin, r, key=it["key"], kind=kind, fmt=fmt)
        # capex module
        cap = sp.get("capex")
        if cap:
            r += 1
            r = self._add(self.in_rows, "section", {"text": f"{n}. Capex allocation — percentages apply to TOTAL project capex (USD m); pre-2024 column = spend before FY2024 (sunk)"}, self.kin, r)
            n += 1
            r = self._add(self.in_rows, "input", {"key": "capex_total_usd", "label": "Total project capex (lifetime investment)", "unit": "USD m", "kind": "scalar",
                                                  "value": cap["total_usd"], "source": cap.get("source", ""), "cls": cap.get("cls", "v29"), "note": cap.get("note", ""), "fmt": "usd_m"},
                          self.kin, r, key="capex_total_usd", kind="scalar", fmt="usd_m")
            pct = {y: cap.get("pct", {}).get(y, 0.0) for y in YEARS}
            r = self._add(self.in_rows, "input", {"key": "capex_pct", "label": "Annual spending % of total capex (2023 column = pre-2024 spend)", "unit": "%", "kind": "annual",
                                                  "values": pct, "source": cap.get("pct_source", ""), "cls": cap.get("pct_cls", "v29"), "fmt": "pct",
                                                  "note": "Move spend between years by editing these percentages; the row must add to 100%. Columns up to the valuation year are pre-valuation estimates unless confirmed below.",
                                                  "total": True}, self.kin, r, key="capex_pct", kind="annual", fmt="pct")
            r = self._add(self.in_rows, "calc", {"key": "capex_usd_view", "label": "Annual capex (USD m) = total x %", "unit": "USD m", "f": "={in:capex_total_usd}*{in:capex_pct}", "fmt": "usd_m", "total": True},
                          self.kin, r, key="capex_usd_view", fmt="usd_m")
            r = self._add(self.in_rows, "calc", {"key": "fx_view", "label": "Exchange rate applied (Global Inputs)", "unit": "IDR/USD", "f": "={g:fx}", "fmt": "fx", "link": True},
                          self.kin, r, key="fx_view", fmt="fx")
            r = self._add(self.in_rows, "calc", {"key": "capex_idr_view", "label": "Annual capex (IDR bn) = USD m x FX / 1,000", "unit": "IDR bn", "f": "={in:capex_usd_view}*{in:fx_view}/1000", "fmt": "idr_bn", "total": True},
                          self.kin, r, key="capex_idr_view", fmt="idr_bn")
            r = self._add(self.in_rows, "scalarcalc", {"key": "capex_pre_val", "label": "Spent up to the valuation date (pre-2024 + FY2024 to valuation year) — estimate unless confirmed", "unit": "USD m",
                                                       "f": "=SUMPRODUCT({in:capex_usd_view@all},--({g:year@all}<={g:val_year}))", "fmt": "usd_m"},
                          self.kin, r, key="capex_pre_val", kind="scalar", fmt="usd_m")
            r = self._add(self.in_rows, "scalarcalc", {"key": "capex_post_val", "label": "Remaining spend after the valuation date (enters the FY2027+ DCF)", "unit": "USD m",
                                                       "f": "=SUMPRODUCT({in:capex_usd_view@all},--({g:year@all}>{g:val_year}))", "fmt": "usd_m"},
                          self.kin, r, key="capex_post_val", kind="scalar", fmt="usd_m")
            r = self._add(self.in_rows, "scalarcalc", {"key": "capex_unalloc", "label": "Unallocated (total - allocated); must be nil", "unit": "USD m",
                                                       "f": "={in:capex_total_usd}-SUM({in:capex_usd_view@all})", "fmt": "usd_m"},
                          self.kin, r, key="capex_unalloc", kind="scalar", fmt="usd_m")
            r = self._add(self.in_rows, "scalarcalc", {"key": "capex_check", "label": "Allocation check (percentages add to 100%)", "unit": "text",
                                                       "f": '=IF(ABS(SUM({in:capex_pct@all})-1)<0.000001,"PASS","CHECK: % not 100%")', "fmt": "text", "check": True},
                          self.kin, r, key="capex_check", kind="scalar", fmt="text")
            cc = cap.get("commissioning_year")
            if cc is not None:
                r = self._add(self.in_rows, "input", {"key": "capex_comm", "label": "Commissioning year of the capex (depreciation starts; may differ from the business start year)", "unit": "year", "kind": "scalar",
                                                      "value": cc, "source": cap.get("commissioning_source", ""), "cls": cap.get("cls", "v29"), "fmt": "year"},
                              self.kin, r, key="capex_comm", kind="scalar", fmt="year")
            elif sp.get("commissioning"):
                r = self._add(self.in_rows, "link", {"label": "Commissioning year of the capex (= project commissioning year)", "unit": "year", "f": "={in:commissioning}", "fmt": "year", "src": "Same as the operating start above"},
                              self.kin, r, key="capex_comm", kind="scalar", fmt="year")
            else:
                r = self._add(self.in_rows, "input", {"key": "capex_comm", "label": "Commissioning year of the capex (depreciation starts)", "unit": "year", "kind": "scalar",
                                                      "value": 2024, "source": "Existing business: first forecast year", "cls": "v29", "fmt": "year"},
                              self.kin, r, key="capex_comm", kind="scalar", fmt="year")
            r = self._add(self.in_rows, "input", {"key": "capex_hist_confirmed", "label": "Historical / pre-valuation spend confirmed from actuals? (1 = confirmed, 0 = v29 projection / unconfirmed estimate)", "unit": "0/1", "kind": "scalar",
                                                  "value": cap.get("hist_confirmed", 0), "source": "User confirmation required", "cls": "PROV", "fmt": "flag01",
                                                  "note": "Pre-valuation capex is carried at v29 projected timing; actual spend to date has not been supplied."},
                          self.kin, r, key="capex_hist_confirmed", kind="scalar", fmt="flag01")
            r = self._add(self.in_rows, "input", {"key": "capex_bshare", "label": "Building / civil share of capex (balance = equipment)", "unit": "%", "kind": "scalar",
                                                  "value": cap["building_share"], "source": cap.get("building_source", ""), "cls": cap.get("cls", "v29"), "fmt": "pct"},
                          self.kin, r, key="capex_bshare", kind="scalar", fmt="pct")
            r = self._add(self.in_rows, "input", {"key": "life_b", "label": "Useful life — building / civil assets", "unit": "years", "kind": "scalar",
                                                  "value": cap["life_b"], "source": cap.get("lives_source", ""), "cls": cap.get("cls", "v29"), "fmt": "int"},
                          self.kin, r, key="life_b", kind="scalar", fmt="int")
            r = self._add(self.in_rows, "input", {"key": "life_e", "label": "Useful life — equipment", "unit": "years", "kind": "scalar",
                                                  "value": cap["life_e"], "source": cap.get("lives_source", ""), "cls": cap.get("cls", "v29"), "fmt": "int"},
                          self.kin, r, key="life_e", kind="scalar", fmt="int")
        self.in_end = r

    # ---- calc sheet
    def _plan_calc(self):
        sp = self.spec
        r = DATA_ROW0
        r = self._add(self.calc_rows, "section", {"text": "A. Operating period"}, self.kcalc, r)
        if sp.get("commissioning"):
            f = "=IF({y}>={in:commissioning},1,0)*{flag}"
        else:
            f = "={flag}"
        r = self._add(self.calc_rows, "row", {"key": "opflag", "label": "Operating flag (1 = project operating in the year)", "unit": "flag", "f": f, "fmt": "flag"},
                      self.kcalc, r, key="opflag", fmt="flag")
        cur = None
        letters = "BCDEFGHIJKLMNOPQRSTUVWXYZ"
        n = 1
        for it in sp.get("calc", []):
            sec = it.get("section", "Calculations")
            if sec != cur:
                cur = sec
                r += 1
                r = self._add(self.calc_rows, "section", {"text": f"{letters[n]}. {sec}"}, self.kcalc, r)
                n += 1
            fmt = it.get("fmt") or S.UNIT_FMT.get(it.get("unit", ""), "num2")
            kind = it.get("kind", "annual")
            r = self._add(self.calc_rows, "row", it, self.kcalc, r, key=it["key"], kind=kind, fmt=fmt)
        # capex & depreciation
        r += 1
        r = self._add(self.calc_rows, "section", {"text": f"{letters[n]}. Capex, depreciation and fixed assets (project)"}, self.kcalc, r)
        n += 1
        cap = sp.get("capex")
        if cap:
            rows = [
                ("capex_usd", "Capex (USD m) — from Input allocation", "USD m", "={in:capex_usd_view}", "usd_m", True),
                ("fx", "Exchange rate (IDR/USD)", "IDR/USD", "={g:fx}", "fx", False),
                ("capex_idr", "Capex (IDR bn)", "IDR bn", "={in:capex_idr_view}", "idr_bn", True),
                ("capex_cum", "Cumulative capex (IDR bn)", "IDR bn", "={c:capex_idr@cum}", "idr_bn", False),
                ("capflag", "Capex in service flag (1 from the capex commissioning year)", "flag", "=IF({y}>={in:capex_comm},1,0)*{flag}", "flag", False),
                ("dep_start", "Depreciation start year of the year's spend = MAX(spend year, capex commissioning year)", "year", "=MAX({y},{in:capex_comm})", "year", False),
                ("dep_b", "Depreciation — building / civil", "IDR bn",
                 "=IF({g:dep_method}=1,SUM({c:capex_idr@all})*{in:capex_bshare}/{in:life_b}*{c:capflag},"
                 "SUMPRODUCT({c:capex_idr@all},--({c:dep_start@all}<={y}),--({y}<{c:dep_start@all}+{in:life_b}))*{in:capex_bshare}/{in:life_b})", "idr_bn", True),
                ("dep_e", "Depreciation — equipment", "IDR bn",
                 "=IF({g:dep_method}=1,SUM({c:capex_idr@all})*(1-{in:capex_bshare})/{in:life_e}*{c:capflag},"
                 "SUMPRODUCT({c:capex_idr@all},--({c:dep_start@all}<={y}),--({y}<{c:dep_start@all}+{in:life_e}))*(1-{in:capex_bshare})/{in:life_e})", "idr_bn", True),
                ("dep_new", "Depreciation on project capex", "IDR bn", "={c:dep_b}+{c:dep_e}", "idr_bn", True),
            ]
            for key, lab, unit, f, fmt, tot in rows:
                r = self._add(self.calc_rows, "row", {"key": key, "label": lab, "unit": unit, "f": f, "fmt": fmt, "total": tot,
                                                      "src": ("Method 1 (v29): total capex x share / life from commissioning, no end. Method 2: each year's spend over its life from MAX(spend year, commissioning)." if key == "dep_b" else None)},
                              self.kcalc, r, key=key, fmt=fmt)
            r = self._add(self.calc_rows, "row", {"key": "nbv_new", "label": "Net book value of project capex (opening = pre-2024 spend)", "unit": "IDR bn",
                                                  "f": "={c:nbv_new@prev}+{c:capex_idr}-{c:dep_new}", "f0": "={c:capex_idr}", "fmt": "idr_bn"},
                          self.kcalc, r, key="nbv_new", fmt="idr_bn")
            r = self._add(self.calc_rows, "row", {"key": "capex_first", "label": "Helper: first year with capex spend (1/0)", "unit": "flag",
                                                  "f": "=IF(AND({c:capex_cum}>0,{c:capex_cum@prev}=0),1,0)", "f0": "=IF({c:capex_idr}>0,1,0)", "fmt": "flag"},
                          self.kcalc, r, key="capex_first", fmt="flag")
        else:
            for key, lab in [("capex_usd", "Capex (USD m) — none in this block"), ("capex_idr", "Capex (IDR bn) — none in this block"),
                             ("dep_new", "Depreciation on project capex — none"), ("nbv_new", "Net book value of project capex — none")]:
                r = self._add(self.calc_rows, "row", {"key": key, "label": lab, "unit": "IDR bn" if "IDR" in lab or "Dep" in lab or "book" in lab else "USD m", "f": "=0",
                                                      "fmt": "idr_bn" if key != "capex_usd" else "usd_m"}, self.kcalc, r, key=key,
                              fmt="idr_bn" if key != "capex_usd" else "usd_m")
        hooks = sp.get("hooks", {})
        da_other = hooks.get("da_other")
        r = self._add(self.calc_rows, "row", {"key": "da_total", "label": "Total depreciation & amortisation (project capex + existing assets / other)", "unit": "IDR bn",
                                              "f": "={c:dep_new}" + (f"+{{c:{da_other}}}" if da_other else ""), "fmt": "idr_bn", "total": True, "bold": True},
                      self.kcalc, r, key="da_total", fmt="idr_bn")
        # working capital
        r += 1
        r = self._add(self.calc_rows, "section", {"text": f"{letters[n]}. Working capital (days from Global Inputs)"}, self.kcalc, r)
        existing = sp.get("existing", False)
        rev, cogs = hooks["revenue"], hooks.get("cogs")
        if existing:
            f0_ar = "={g:open_ar}*{c:" + rev + "@2024}/({s:global:rev2024_existing})"
            f0_ap = ("={g:open_ap}*{c:" + cogs + "@2024}/({s:global:cogs2024_existing})") if cogs else "=0"
            note_ar = "2023 = allocated share of the company Dec-23 receivables (pro rata FY2024 revenue) — allocation, not source data"
        else:
            f0_ar, f0_ap, note_ar = "=0", "=0", None
        r = self._add(self.calc_rows, "row", {"key": "ar", "label": "Trade receivables (revenue x AR days / days)", "unit": "IDR bn",
                                              "f": "={c:" + rev + "}*{g:ar_days}/{days}", "f0": f0_ar, "fmt": "idr_bn", "note": note_ar}, self.kcalc, r, key="ar", fmt="idr_bn")
        r = self._add(self.calc_rows, "row", {"key": "ap", "label": "Trade payables (cost of sales x AP days / days)", "unit": "IDR bn",
                                              "f": ("={c:" + cogs + "}*{g:ap_days}/{days}") if cogs else "=0", "f0": f0_ap, "fmt": "idr_bn"}, self.kcalc, r, key="ap", fmt="idr_bn")
        r = self._add(self.calc_rows, "row", {"key": "nwc", "label": "Net working capital (receivables - payables)", "unit": "IDR bn", "f": "={c:ar}-{c:ap}", "fmt": "idr_bn"},
                      self.kcalc, r, key="nwc", fmt="idr_bn")
        r = self._add(self.calc_rows, "row", {"key": "dnwc", "label": "Increase in net working capital (cash outflow +)", "unit": "IDR bn", "f": "={c:nwc}-{c:nwc@prev}", "f0": "=0", "fmt": "idr_bn", "total": True},
                      self.kcalc, r, key="dnwc", fmt="idr_bn")
        self.calc_end = r

    # ---- output sheet
    def _plan_output(self):
        sp = self.spec
        hooks = sp.get("hooks", {})
        rev, cogs, opex = hooks["revenue"], hooks.get("cogs"), hooks.get("opex")
        oth = hooks.get("other_income")
        r = DATA_ROW0
        r = self._add(self.out_rows, "section", {"text": "A. Income statement — project, unlevered (before corporate financing; tax on project EBIT)"}, self.kout, r)
        rows = [
            ("revenue", "Revenue", "={c:" + rev + "}", False, "idr_bn"),
            ("cogs", "Cost of sales", ("=-{c:" + cogs + "}") if cogs else "=0", False, "idr_bn"),
            ("gp", "Gross profit", "={o:revenue}+{o:cogs}", True, "idr_bn"),
            ("opex", "Operating expenses", ("=-{c:" + opex + "}") if opex else "=0", False, "idr_bn"),
        ]
        if oth:
            rows.append(("other_income", "Other operating income", "={c:" + oth + "}", False, "idr_bn"))
        rows += [
            ("ebitda", "EBITDA", "={o:gp}+{o:opex}" + ("+{o:other_income}" if oth else ""), True, "idr_bn"),
            ("da", "Depreciation & amortisation", "=-{c:da_total}", False, "idr_bn"),
            ("ebit", "EBIT", "={o:ebitda}+{o:da}", True, "idr_bn"),
            ("tax", "Tax on project EBIT (rate x MAX(0, EBIT); v29 line method — no loss relief between lines)", "=-MAX(0,{o:ebit})*{g:tax_rate}", False, "idr_bn"),
            ("nopat", "Net operating profit after tax (NOPAT)", "={o:ebit}+{o:tax}", True, "idr_bn"),
        ]
        for key, lab, f, bold, fmt in rows:
            r = self._add(self.out_rows, "row", {"key": key, "label": lab, "unit": "IDR bn", "f": f, "fmt": fmt, "bold": bold, "total": True}, self.kout, r, key=key, fmt=fmt)
        r = self._add(self.out_rows, "row", {"key": "m_ebitda", "label": "EBITDA margin", "unit": "%", "f": "=IF({o:revenue}=0,0,{o:ebitda}/{o:revenue})", "fmt": "pct"}, self.kout, r, key="m_ebitda", fmt="pct")
        r = self._add(self.out_rows, "row", {"key": "m_ebit", "label": "EBIT margin", "unit": "%", "f": "=IF({o:revenue}=0,0,{o:ebit}/{o:revenue})", "fmt": "pct"}, self.kout, r, key="m_ebit", fmt="pct")
        r += 1
        r = self._add(self.out_rows, "section", {"text": "B. Free cash flow to firm (unlevered)"}, self.kout, r)
        rows = [
            ("f_nopat", "NOPAT", "={o:nopat}", False),
            ("f_da", "Add back depreciation & amortisation", "=-{o:da}", False),
            ("f_capex", "Less capex (IDR bn)", "=-{c:capex_idr}", False),
            ("f_dnwc", "Less increase in net working capital", "=-{c:dnwc}", False),
            ("fcff", "Free cash flow to firm (FCFF)", "={o:f_nopat}+{o:f_da}+{o:f_capex}+{o:f_dnwc}", True),
            ("fcff_cum", "Cumulative FCFF", "={o:fcff@cum}", False),
            ("fcff_dcf", "FCFF in the company DCF window (years after the valuation date)", "={o:fcff}*{g:dcf_flag}", False),
            ("fund_need", "Funding need = cash shortfall to be funded by FSL (MAX(0, -FCFF))", "=MAX(0,-{o:fcff})", False),
        ]
        for key, lab, f, bold in rows:
            r = self._add(self.out_rows, "row", {"key": key, "label": lab, "unit": "IDR bn", "f": f, "fmt": "idr_bn", "bold": bold, "total": key not in ("fcff_cum",)}, self.kout, r, key=key, fmt="idr_bn")
        r += 1
        r = self._add(self.out_rows, "section", {"text": "C. Project balance sheet — net operating assets (no project-level cash or debt: funded by FSL corporate)"}, self.kout, r)
        rows = [
            ("bs_capex_usd", "Capex in the year (USD m)", "={c:capex_usd}", "usd_m", False),
            ("bs_capex", "Capex in the year (IDR bn)", "={c:capex_idr}", "idr_bn", False),
            ("bs_ppe", "Fixed assets — project capex at net book value", "={c:nbv_new}", "idr_bn", False),
            ("bs_ar", "Trade receivables", "={c:ar}", "idr_bn", False),
            ("bs_ap", "Trade payables", "=-{c:ap}", "idr_bn", False),
            ("bs_noa", "Net operating assets funded by FSL", "={o:bs_ppe}+{o:bs_ar}+{o:bs_ap}", "idr_bn", True),
        ]
        for key, lab, f, fmt, bold in rows:
            r = self._add(self.out_rows, "row", {"key": key, "label": lab, "unit": "USD m" if fmt == "usd_m" else "IDR bn", "f": f, "fmt": fmt, "bold": bold,
                                                 "total": key in ("bs_capex_usd", "bs_capex")}, self.kout, r, key=key, fmt=fmt)
        if sp.get("existing"):
            r = self._add(self.out_rows, "note", {"text": "Existing fixed assets of this business are held at company level (Dec-23 net book value not allocated by business in v29); their depreciation is charged in this project's P&L."}, self.kout, r)
        r += 1
        r = self._add(self.out_rows, "section", {"text": "D. Investment returns"}, self.kout, r)
        # helper rows for payback
        r = self._add(self.out_rows, "row", {"key": "pb_cum_prev", "label": "Payback helper: payback already reached in an earlier year (1/0)", "unit": "flag",
                                             "f": "=MIN(1,{o:pb_cum_prev@prev}+{o:pb_flag@prev})", "f0": "=0", "fmt": "flag"}, self.kout, r, key="pb_cum_prev", fmt="flag")
        r = self._add(self.out_rows, "row", {"key": "pb_flag", "label": "Payback helper: first year in which cumulative FCFF turns non-negative after being negative (1/0)", "unit": "flag",
                                             "f": "=IF(AND({o:fcff_cum}>=0,{o:fcff_cum@prev}<0,{o:pb_cum_prev}=0),1,0)", "f0": "=0", "fmt": "flag"}, self.kout, r, key="pb_flag", fmt="flag")
        r = self._add(self.out_rows, "row", {"key": "tv_flow", "label": "Terminal value at end-2036 (Gordon growth on FY2036 FCFF at WACC and g from Global Inputs) — in 2036 column", "unit": "IDR bn",
                                             "f": "=IF({y}={g:last_year},{o:fcff}*(1+{g:tg})/({g:wacc}-{g:tg}),0)", "fmt": "idr_bn"}, self.kout, r, key="tv_flow", fmt="idr_bn")
        r = self._add(self.out_rows, "row", {"key": "fcff_tv", "label": "FCFF incl. terminal value", "unit": "IDR bn", "f": "={o:fcff}+{o:tv_flow}", "fmt": "idr_bn"}, self.kout, r, key="fcff_tv", fmt="idr_bn")
        r = self._add(self.out_rows, "row", {"key": "pv", "label": "Present value at the valuation date of FCFF incl. terminal value (years after the valuation date only)", "unit": "IDR bn",
                                             "f": "={o:fcff_tv}*{g:df}", "fmt": "idr_bn", "total": True}, self.kout, r, key="pv", fmt="idr_bn")
        existing = sp.get("existing", False)
        scal = [
            ("r_npv", "Value at the valuation date (31-Dec-2026): PV of FY2027+ FCFF incl. terminal value, at WACC", "IDR bn", "=SUM({o:pv@all})", "idr_bn"),
            ("r_npv_notv", "  of which PV of FY2027-FY2036 FCFF (excl. terminal value)", "IDR bn", "=SUMPRODUCT({o:fcff@all},{g:df@all})", "idr_bn"),
            ("r_irr", "Full-life project IRR — FCFF 2023-2036 incl. pre-2024 spend, no terminal value" if not existing else "Full-life project IRR",
             "%", ('=IFERROR(IRR({o:fcff@all}),"n/a")' if not existing else '="n/a — existing business (no entry cost / original investment data)"'), "pct"),
            ("r_irr_tv", "Full-life project IRR incl. terminal value at end-2036" if not existing else "Full-life project IRR incl. terminal value",
             "%", ('=IFERROR(IRR({o:fcff_tv@all}),"n/a")' if not existing else '="n/a — existing business"'), "pct"),
            ("r_eirr", "Equity IRR", "%", ('="see section D2 (project debt)"' if sp.get("project_debt") else ('="= project IRR (no project-level debt in v29; corporate LT loan is not project debt)"' if not existing else '="n/a — existing business"')), "pct"),
            ("r_pb_year", "Payback year (cumulative FCFF turns positive)", "year", '=IF(SUM({o:pb_flag@all})=0,"not within horizon",SUMPRODUCT({o:pb_flag@all},{g:year@all}))', "year"),
            ("r_first_spend", "First year of capex spend (2023 = pre-2024)", "year", ('=IF(SUM({c:capex_usd@all})=0,"n/a",SUMPRODUCT({c:capex_first@all},{g:year@all}))' if sp.get("capex") else '="n/a"'), "year"),
            ("r_capex_total", "Total capex (USD m) / of which after the valuation date", "USD m", "={in:capex_total_usd}" if sp.get("capex") else "=0", "usd_m"),
            ("r_capex_post", "Capex after the valuation date (USD m)", "USD m", "={in:capex_post_val}" if sp.get("capex") else "=0", "usd_m"),
        ]
        for key, lab, unit, f, fmt in scal:
            r = self._add(self.out_rows, "scalarcalc", {"key": key, "label": lab, "unit": unit, "f": f, "fmt": fmt}, self.kout, r, key=key, kind="scalar", fmt=fmt)
        pdebt = sp.get("project_debt")
        if pdebt:
            r += 1
            r = self._add(self.out_rows, "section", {"text": "D2. Equity cash flow with project debt (project-level financing from the Calc sheet) — equity IRR"}, self.kout, r)
            rows = [
                ("e_fcff", "FCFF", "={o:fcff}", False),
                ("e_draw", "+ Project debt drawdown", "={c:" + pdebt["draw"] + "}", False),
                ("e_repay", "- Project debt repayment", "=-{c:" + pdebt["repay"] + "}", False),
                ("e_int", "- Interest after tax shield (interest x (1 - tax rate))", "=-{c:" + pdebt["interest"] + "}*(1-{g:tax_rate})", False),
                ("fcfe", "Free cash flow to equity (100% of the project)", "={o:e_fcff}+{o:e_draw}+{o:e_repay}+{o:e_int}", True),
                ("fcfe_fsl", "Free cash flow to FSL = FCFE x FSL economic interest", "={o:fcfe}*{in:eint}", False),
                ("fcfe_tv", "FCFE incl. terminal value (equity value at end-2036 = FCFF terminal value less debt outstanding)", "={o:fcfe}+IF({y}={g:last_year},{o:tv_flow}-{c:" + pdebt["balance"] + "},0)", False),
            ]
            for key, lab, f, bold in rows:
                r = self._add(self.out_rows, "row", {"key": key, "label": lab, "unit": "IDR bn", "f": f, "fmt": "idr_bn", "bold": bold, "total": True}, self.kout, r, key=key, fmt="idr_bn")
            r = self._add(self.out_rows, "scalarcalc", {"key": "r_eirr2", "label": "Equity IRR — FCFE 2023-2036, no terminal value", "unit": "%", "f": '=IFERROR(IRR({o:fcfe@all}),"n/a")', "fmt": "pct"}, self.kout, r, key="r_eirr2", kind="scalar", fmt="pct")
            r = self._add(self.out_rows, "scalarcalc", {"key": "r_eirr2_tv", "label": "Equity IRR — FCFE incl. terminal value", "unit": "%", "f": '=IFERROR(IRR({o:fcfe_tv@all}),"n/a")', "fmt": "pct"}, self.kout, r, key="r_eirr2_tv", kind="scalar", fmt="pct")
            r = self._add(self.out_rows, "note", {"text": "Project debt is a project-level financing assumption (DUMMY for Badas); it is not part of the corporate LT loan and is not added to the company balance sheet while the project is outside the valuation."}, self.kout, r)
        r += 1
        r = self._add(self.out_rows, "section", {"text": "E. Case inclusion (drives the company statements and valuations)"}, self.kout, r)
        r = self._add(self.out_rows, "scalarcalc", {"key": "flag_wo", "label": "Included in the Without-BUMN case (1/0)", "unit": "0/1", "f": "={in:incl}*IF({in:case}=1,1,0)", "fmt": "flag01"}, self.kout, r, key="flag_wo", kind="scalar", fmt="flag01")
        r = self._add(self.out_rows, "scalarcalc", {"key": "flag_w", "label": "Included in the With-BUMN case (1/0)", "unit": "0/1", "f": "={in:incl}", "fmt": "flag01"}, self.kout, r, key="flag_w", kind="scalar", fmt="flag01")
        r = self._add(self.out_rows, "scalarcalc", {"key": "eint", "label": "FSL economic interest", "unit": "%", "f": "={in:eint}", "fmt": "pct"}, self.kout, r, key="eint", kind="scalar", fmt="pct")
        self.out_end = r


# ----------------------------------------------------------------------------------------------- writing
def write_project(model: Model, plan: ProjectPlan):
    sp = plan.spec
    pid = plan.pid
    t_in, t_calc, t_out = sheet_titles(sp)
    subtitle = f"{sp['name']} — {sp.get('entity','')} — {sp.get('category','')}. IDR bn unless stated. Blue = input; black = formula; green = link."
    extra = [("Input", t_in), ("Calc", t_calc), ("Output", t_out)]
    # --- Input sheet
    model.write_ts_header(plan.kin, f"{sp['id']} {sp['name']} — " + ("FS WORKING PAPER (INPUT)" if sp.get("working_paper") else "INPUT"), subtitle)
    model.nav_links(plan.kin, row=3, extra=extra)
    if sp.get("working_paper"):
        wsi = model.ws(plan.kin)
        for col, h, w in ((22, "Definition (what the FS should provide)", 52), (23, "Period", 16), (24, "Source / status", 30), (25, "Calculation destination", 34), (26, "Question for the operations team", 52)):
            c = wsi.cell(5, col)
            c.value = h
            c.font = S.F_HDR
            c.fill = S.FILL_HDR
            wsi.column_dimensions[L(col)].width = w
    ctx = {"sheet": plan.kin, "pid": pid}
    ws = model.ws(plan.kin)
    for typ, it, r in plan.in_rows:
        _write_row(model, plan.kin, typ, it, r, ctx, is_input_sheet=True)
    # --- Calc sheet
    model.write_ts_header(plan.kcalc, f"{sp['id']} {sp['name']} — CALCULATION", "Drivers and supporting schedules. All cells are formulas; edit assumptions on the Input sheet.")
    model.nav_links(plan.kcalc, row=3, extra=extra)
    ctx = {"sheet": plan.kcalc, "pid": pid}
    for typ, it, r in plan.calc_rows:
        _write_row(model, plan.kcalc, typ, it, r, ctx)
    # --- Output sheet
    model.write_ts_header(plan.kout, f"{sp['id']} {sp['name']} — OUTPUT", "Project income statement, FCFF, net operating assets, funding need and returns (unlevered, IDR bn). Linked into Project Summary, FSL Financials and Valuation.")
    model.nav_links(plan.kout, row=3, extra=extra)
    ctx = {"sheet": plan.kout, "pid": pid}
    for typ, it, r in plan.out_rows:
        _write_row(model, plan.kout, typ, it, r, ctx)
    for k in (plan.kin, plan.kcalc, plan.kout):
        model.set_print(k)


def _write_row(model: Model, sheetkey, typ, it, r, ctx, is_input_sheet=False):
    ws = model.ws(sheetkey)
    if typ == "section":
        model.section(sheetkey, r, it["text"])
        return
    if typ == "note":
        ws.cell(r, LABEL_COL).value = it["text"]
        ws.cell(r, LABEL_COL).font = S.F_NOTE
        return
    if typ == "text":
        model.write_label(ws, r, it["label"])
        c = ws.cell(r, SCALAR_COL)
        c.value = it["value"]
        c.font = S.F_BODY
        c.alignment = S.A_LEFT
        return
    fmt = it.get("fmt") or S.UNIT_FMT.get(it.get("unit", ""), "num2")
    if typ == "input" and fmt == "pct":
        # very small rates (e.g. insurance 0.015% of sales) would display as 0.00% -> use a finer format automatically
        vals = [v for v in (it.get("values", {}).values() if it.get("kind", "annual") == "annual" else [it.get("value")]) if isinstance(v, (int, float))]
        if vals and 0 < max(abs(v) for v in vals) < 0.001:
            fmt = "pct6" if max(abs(v) for v in vals) < 0.0001 else "pct4"
    if typ == "link":
        model.write_label(ws, r, it["label"], unit=it.get("unit"), typ="link", src=it.get("src"), note=it.get("note"))
        model.write_scalar(sheetkey, r, it["f"], fmt, ctx=ctx, font=S.F_LINK)
        return
    if typ == "input":
        cls = it.get("cls", "v29")
        kind = it.get("kind", "annual")
        label = it["label"]
        if cls == "DUMMY":
            label = label + "   [DUMMY — REPLACE WITH FS DATA]"
        model.write_label(ws, r, label, unit=it.get("unit"), typ=cls, src=it.get("source"), note=it.get("note"))
        for col, k in ((22, "definition"), (23, "period"), (24, "status"), (25, "dest"), (26, "question")):
            if it.get(k):
                c = ws.cell(r, col)
                c.value = it[k]
                c.font = S.font(size=8, color="C00000" if (k == "status" and cls == "DUMMY") else "404040", bold=(k == "status"))
                c.alignment = S.A_WRAP if k in ("definition", "question") else S.A_LEFT
        ws.cell(r, TYPE_COL).font = S.font(bold=(cls in ("DUMMY", "CHANGE", "UPDATED", "PROV")), color="C00000" if cls == "DUMMY" else "595959", size=8)
        fill = CLASS_FILL.get(cls, S.FILL_INPUT)
        if kind == "scalar":
            v = it.get("value")
            model.write_scalar(sheetkey, r, v, fmt, ctx=ctx, font=S.F_INPUT if not (isinstance(v, str) and v.startswith("=")) else S.F_LINK, fill=fill)
        else:
            vals = it.get("values", {})
            model.write_annual_values(sheetkey, r, vals, fmt, fill=fill)
            if it.get("total"):
                t = ws.cell(r, SCALAR_COL)
                t.value = f"=SUM({FCOL}{r}:{LCOL}{r})"
                t.number_format = S.NF[fmt]
                t.font = S.F_BOLD
                t.alignment = S.A_RIGHT
        if cls == "DUMMY":
            for c in range(LABEL_COL, LAST_COL + 1):
                if ws.cell(r, c).fill is None or ws.cell(r, c).fill.fgColor.rgb in (None, "00000000"):
                    ws.cell(r, c).fill = S.FILL_DUMMY
            ws.cell(r, LABEL_COL).fill = S.FILL_DUMMY
            ws.cell(r, LABEL_COL).font = S.font(bold=True, color="C00000")
        return
    if typ in ("row", "calc"):
        kind = it.get("kind", "annual")
        font = S.F_LINK if it.get("link") else None
        model.write_label(ws, r, it["label"], unit=it.get("unit"), typ=it.get("typ"), bold=it.get("bold", False), src=it.get("src"), note=it.get("note"),
                          indent=it.get("indent", 0))
        if kind == "scalar":
            model.write_scalar(sheetkey, r, it["f"], fmt, ctx=ctx, font=font or (S.F_BOLD if it.get("bold") else S.F_BODY))
        else:
            model.write_annual_formula(sheetkey, r, it["f"], ctx, fmt, f0=it.get("f0"), years=it.get("years"), font=font,
                                       total=it.get("total", False), bold=it.get("bold", False),
                                       border=S.B_TOTAL if it.get("bold") else None)
        return
    if typ == "scalarcalc":
        model.write_label(ws, r, it["label"], unit=it.get("unit"), src=it.get("src"), note=it.get("note"), bold=it.get("bold", False))
        cell = model.write_scalar(sheetkey, r, it["f"], fmt, ctx=ctx, font=S.F_BOLD if it.get("bold") else S.F_BODY)
        if it.get("check"):
            cell.alignment = S.A_CENTER
        return
    raise ValueError(typ)
