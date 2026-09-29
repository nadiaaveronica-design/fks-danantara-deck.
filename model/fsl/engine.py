"""Core build engine: timeline, row registry, formula-template resolution, sheet helpers.

Every time-series sheet uses the same column layout so cross-sheet references keep the
column letter:  A spacer | B label | C unit | D type/flag | E scalar-or-total | F..S = 2023..2036 | T source | U note
Header rows on time-series sheets: 1 title, 2 subtitle, 3 navigation, 5 Year, 6 Period end, 7 Days, 8 Model flag.
Data starts at row 10.
"""
import re
from openpyxl import Workbook
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.hyperlink import Hyperlink
from . import styles as S

YEARS = list(range(2023, 2037))          # 14 columns: F (2023, opening / pre-2024) .. S (2036)
FIRST_COL = 6                             # F
LAST_COL = FIRST_COL + len(YEARS) - 1     # S
YEAR_COL = {y: FIRST_COL + i for i, y in enumerate(YEARS)}
COL_YEAR = {c: y for y, c in YEAR_COL.items()}
FCOL = get_column_letter(FIRST_COL)
LCOL = get_column_letter(LAST_COL)
FORECAST_YEARS = YEARS[1:]                # 2024..2036
ROW_YEAR, ROW_PEND, ROW_DAYS, ROW_FLAG = 5, 6, 7, 8
DATA_ROW0 = 10
SCALAR_COL = 5                            # E
SRC_COL, NOTE_COL = 20, 21                # T, U
LABEL_COL, UNIT_COL, TYPE_COL = 2, 3, 4

GLOBAL_TITLE = "Global Inputs"


def L(c):
    return get_column_letter(c)


def yc(year):
    """Column letter of a model year."""
    return L(YEAR_COL[year])


class RowRef:
    __slots__ = ("sheet", "row", "kind", "fmt", "label", "col")

    def __init__(self, sheet, row, kind, fmt=None, label="", col=None):
        self.sheet, self.row, self.kind, self.fmt, self.label = sheet, row, kind, fmt, label
        self.col = col if col is not None else SCALAR_COL


class Model:
    """Workbook wrapper holding the registry of named rows across sheets."""

    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.sheets = {}      # sheetkey -> worksheet
        self.titles = {}      # sheetkey -> title
        self.reg = {}         # (sheetkey, key) -> RowRef
        self.projects = {}    # pid -> spec
        self.order = []       # sheet keys in creation order
        self.notes = []       # (sheetkey, text) for later documentation

    # ---------------------------------------------------------------- sheets
    def add_sheet(self, key, title, tab_color=None, ts=True, widths=None):
        if len(title) > 31:
            raise ValueError(f"sheet title too long: {title}")
        ws = self.wb.create_sheet(title)
        self.sheets[key] = ws
        self.titles[key] = title
        self.order.append(key)
        if tab_color:
            ws.sheet_properties.tabColor = tab_color
        ws.sheet_view.showGridLines = False
        if ts:
            ws.column_dimensions["A"].width = 2
            ws.column_dimensions["B"].width = 66
            ws.column_dimensions["C"].width = 11
            ws.column_dimensions["D"].width = 12
            ws.column_dimensions["E"].width = 13
            for c in range(FIRST_COL, LAST_COL + 1):
                ws.column_dimensions[L(c)].width = 11.5
            ws.column_dimensions[L(SRC_COL)].width = 40
            ws.column_dimensions[L(NOTE_COL)].width = 48
        if widths:
            for col, w in widths.items():
                ws.column_dimensions[col].width = w
        return ws

    def ws(self, key):
        return self.sheets[key]

    def title(self, key):
        return self.titles[key]

    # -------------------------------------------------------------- registry
    def register(self, sheetkey, key, row, kind="annual", fmt=None, label="", col=None):
        k = (sheetkey, key)
        if k in self.reg:
            raise KeyError(f"duplicate key {key} on sheet {sheetkey}")
        self.reg[k] = RowRef(sheetkey, row, kind, fmt, label, col)
        return row

    def get(self, sheetkey, key):
        try:
            return self.reg[(sheetkey, key)]
        except KeyError:
            raise KeyError(f"unknown key '{key}' on sheet '{sheetkey}' (known: "
                           f"{sorted(k for s, k in self.reg if s == sheetkey)[:40]} ...)")

    def has(self, sheetkey, key):
        return (sheetkey, key) in self.reg

    def cell(self, sheetkey, key, col=None, year=None, absolute=False, cur_sheet=None):
        """Address of a registered row in a given column (annual) or its scalar cell."""
        r = self.get(sheetkey, key)
        if r.kind == "scalar":
            addr = f"${L(r.col)}${r.row}"
        else:
            if year is not None:
                col = YEAR_COL[year]
                addr = f"${L(col)}${r.row}"
            elif absolute:
                addr = f"${L(col)}${r.row}"
            else:
                addr = f"{L(col)}{r.row}"
        if cur_sheet == sheetkey:
            return addr
        return f"'{self.titles[sheetkey]}'!{addr}"

    def rng(self, sheetkey, key, cur_sheet=None, years=None):
        r = self.get(sheetkey, key)
        c1, c2 = (FIRST_COL, LAST_COL) if years is None else (YEAR_COL[years[0]], YEAR_COL[years[1]])
        addr = f"${L(c1)}${r.row}:${L(c2)}${r.row}"
        if cur_sheet == sheetkey:
            return addr
        return f"'{self.titles[sheetkey]}'!{addr}"

    # ------------------------------------------------------- template resolve
    _TOKEN = re.compile(r"\{([a-zA-Z]+):([^{}]+)\}|\{(y|yprev|days|col|prevcol|flag|ycol)\}")

    def resolve(self, template, col, ctx):
        """Resolve a formula template for a given column index.

        ctx: dict with 'sheet' (current sheetkey), 'pid' (project id), optional 'aliases'.
        Tokens:
          {in:key} {in:key@2025}          project input row (annual same column / fixed year / scalar)
          {c:key} {c:key@prev} {c:key@2025} {c:key@all} {c:key@cum}    project calc row
          {o:key}                         project output row
          {x:PID:key}  {x:PID:key@prev}   another project's calc (or output 'out.' key) row
          {xin:PID:key}                   another project's input row
          {g:key} {g:key@2026} {g:key@all} {g:key@prev}   Global Inputs row
          {s:SHEETKEY:key}                any registered sheet key
          {y} {yprev} {days} {flag}       header cells of the current sheet (Year, prior Year, Days, model flag)
          {col} {prevcol}                 column letters
        """
        sheet = ctx["sheet"]
        pid = ctx.get("pid")

        def rowref(sheetkey, key, mod):
            r = self.get(sheetkey, key)
            if r.kind == "scalar":
                return self.cell(sheetkey, key, cur_sheet=sheet)
            if mod is None:
                return self.cell(sheetkey, key, col=col, cur_sheet=sheet)
            if mod == "prev":
                return self.cell(sheetkey, key, col=col - 1, cur_sheet=sheet)
            if mod == "next":
                return self.cell(sheetkey, key, col=col + 1, cur_sheet=sheet)
            if mod == "all":
                return self.rng(sheetkey, key, cur_sheet=sheet)
            if mod == "fc":   # forecast years 2024..2036
                return self.rng(sheetkey, key, cur_sheet=sheet, years=(2024, 2036))
            if mod == "cum":  # cumulative to date SUM($F r : col r)
                rr = r.row
                pre = "" if sheet == sheetkey else f"'{self.titles[sheetkey]}'!"
                return f"SUM({pre}${FCOL}${rr}:{pre}{L(col)}{rr})"
            if mod == "abs":
                return self.cell(sheetkey, key, col=col, absolute=True, cur_sheet=sheet)
            if mod == "e":    # the scalar / flag cell (column E) of an annual row
                pre = "" if sheet == sheetkey else f"'{self.titles[sheetkey]}'!"
                return f"{pre}${L(SCALAR_COL)}${r.row}"
            if mod.isdigit():
                return self.cell(sheetkey, key, year=int(mod), cur_sheet=sheet)
            raise ValueError(f"bad modifier {mod} in token for {key}")

        def sub(m):
            if m.group(3):
                t = m.group(3)
                if t == "y":
                    return f"{L(col)}${ROW_YEAR}"
                if t == "yprev":
                    return f"{L(col-1)}${ROW_YEAR}"
                if t == "days":
                    return f"{L(col)}${ROW_DAYS}"
                if t == "flag":
                    return f"{L(col)}${ROW_FLAG}"
                if t == "col":
                    return L(col)
                if t == "prevcol":
                    return L(col - 1)
                if t == "ycol":
                    return f"${L(col)}${ROW_YEAR}"
            kind, body = m.group(1), m.group(2)
            mod = None
            if "@" in body:
                body, mod = body.split("@", 1)
            if kind == "in":
                return rowref(f"{pid}:in", body, mod)
            if kind == "c":
                return rowref(f"{pid}:calc", body, mod)
            if kind == "o":
                return rowref(f"{pid}:out", body, mod)
            if kind == "g":
                return rowref("global", body, mod)
            if kind in ("x", "xin", "xo"):
                opid, key = body.split(":", 1)
                if kind == "xin":
                    return rowref(f"{opid}:in", key, mod)
                if kind == "xo" or key.startswith("out."):
                    return rowref(f"{opid}:out", key, mod)
                return rowref(f"{opid}:calc", key, mod)
            if kind == "s":
                sk, key = body.split(":", 1)
                return rowref(sk, key, mod)
            if kind == "t":   # this sheet
                return rowref(sheet, body, mod)
            raise ValueError(f"unknown token kind {kind}")

        return self._TOKEN.sub(sub, template)

    # ------------------------------------------------------------ writing
    def write_ts_header(self, sheetkey, title, subtitle, nav=True, freeze=True, pre_shade=True):
        ws = self.ws(sheetkey)
        ws["B1"] = title
        ws["B1"].font = S.F_TITLE
        ws["B2"] = subtitle
        ws["B2"].font = S.F_SUB
        ws.row_dimensions[1].height = 22
        # header rows link to Global Inputs timeline
        ws[f"B{ROW_YEAR}"] = "Year"
        ws[f"B{ROW_PEND}"] = "Period end"
        ws[f"B{ROW_DAYS}"] = "Days in year"
        ws[f"B{ROW_FLAG}"] = "Forecast flag (1 = forecast year)"
        for r in (ROW_YEAR, ROW_PEND, ROW_DAYS, ROW_FLAG):
            ws[f"B{r}"].font = S.F_BOLD if r == ROW_YEAR else S.F_BODY
        for c in range(FIRST_COL, LAST_COL + 1):
            col = L(c)
            if sheetkey == "global":
                ws[f"{col}{ROW_YEAR}"] = COL_YEAR[c]
                ws[f"{col}{ROW_PEND}"] = f"=DATE({col}{ROW_YEAR},12,31)"
                ws[f"{col}{ROW_DAYS}"] = f"={col}{ROW_PEND}-DATE({col}{ROW_YEAR},1,1)+1"
                ws[f"{col}{ROW_FLAG}"] = 1 if COL_YEAR[c] >= 2024 else 0
                ws[f"{col}{ROW_YEAR}"].font = S.F_HDR
            else:
                g = f"'{GLOBAL_TITLE}'!"
                ws[f"{col}{ROW_YEAR}"] = f"={g}{col}${ROW_YEAR}"
                ws[f"{col}{ROW_PEND}"] = f"={g}{col}${ROW_PEND}"
                ws[f"{col}{ROW_DAYS}"] = f"={g}{col}${ROW_DAYS}"
                ws[f"{col}{ROW_FLAG}"] = f"={g}{col}${ROW_FLAG}"
                ws[f"{col}{ROW_YEAR}"].font = S.F_HDR
            ws[f"{col}{ROW_YEAR}"].fill = S.FILL_HDR
            ws[f"{col}{ROW_YEAR}"].alignment = S.A_CENTER
            ws[f"{col}{ROW_YEAR}"].number_format = S.NF["year"]
            ws[f"{col}{ROW_PEND}"].number_format = S.NF["date"]
            ws[f"{col}{ROW_PEND}"].font = S.F_NOTE
            ws[f"{col}{ROW_PEND}"].alignment = S.A_CENTER
            ws[f"{col}{ROW_DAYS}"].number_format = S.NF["days"]
            ws[f"{col}{ROW_DAYS}"].font = S.F_NOTE
            ws[f"{col}{ROW_DAYS}"].alignment = S.A_CENTER
            ws[f"{col}{ROW_FLAG}"].number_format = S.NF["flag01"]
            ws[f"{col}{ROW_FLAG}"].font = S.F_NOTE
            ws[f"{col}{ROW_FLAG}"].alignment = S.A_CENTER
        for c in (LABEL_COL, UNIT_COL, TYPE_COL, SCALAR_COL):
            ws.cell(ROW_YEAR, c).fill = S.FILL_HDR
            ws.cell(ROW_YEAR, c).font = S.F_HDR
        ws.cell(ROW_YEAR, UNIT_COL).value = "Unit"
        ws.cell(ROW_YEAR, TYPE_COL).value = "Type"
        ws.cell(ROW_YEAR, SCALAR_COL).value = "Value / total"
        ws.cell(ROW_YEAR, SRC_COL).value = "Source (v29 reference / basis)"
        ws.cell(ROW_YEAR, NOTE_COL).value = "Note"
        for c in (SRC_COL, NOTE_COL):
            ws.cell(ROW_YEAR, c).fill = S.FILL_HDR
            ws.cell(ROW_YEAR, c).font = S.F_HDR
        if freeze:
            ws.freeze_panes = f"{FCOL}{DATA_ROW0}"
        if nav:
            self.nav_links(sheetkey, row=3)

    def nav_links(self, sheetkey, row=3, extra=None):
        """Standard navigation line: Navigation | Global Inputs | Project Summary | Funding | Financials | Valuation."""
        ws = self.ws(sheetkey)
        ws.cell(row, 2).value = "Go to:"
        ws.cell(row, 2).font = S.F_NOTE
        targets = [("Navigation", "Navigation"), ("Global Inputs", GLOBAL_TITLE), ("Project Summary", "Project Summary"),
                   ("Funding", "Funding Requirement"), ("FSL w/o BUMN", "FSL Financials — Without BUMN"),
                   ("FSL w/ BUMN", "FSL Financials — With BUMN"), ("Val w/o BUMN", "Valuation — Without BUMN"),
                   ("Val w/ BUMN", "Valuation — With BUMN"), ("Checks", "Checks"), ("Model Flow", "Model Flow")]
        if extra:
            targets = extra + targets
        c = 3
        for label, target in targets:
            cell = ws.cell(row, c)
            cell.value = label
            cell.hyperlink = f"#'{target}'!A1"
            cell.font = S.F_NAV
            cell.alignment = S.A_CENTER
            c += 1

    def link(self, ws, cell, target_sheet, label, target_cell="A1"):
        ws[cell] = label
        ws[cell].hyperlink = f"#'{target_sheet}'!{target_cell}"
        ws[cell].font = S.F_NAV

    def section(self, sheetkey, row, text, level=1, note=None):
        ws = self.ws(sheetkey)
        ws.cell(row, 2).value = text
        fill = S.FILL_SECTION if level == 1 else S.FILL_SUBSECTION
        for c in range(2, LAST_COL + 1):
            ws.cell(row, c).fill = fill
        ws.cell(row, 2).font = S.F_BOLD if level == 1 else S.font(bold=True, color="1F3864")
        if note:
            ws.cell(row, SRC_COL).value = note
            ws.cell(row, SRC_COL).font = S.F_NOTE
        return row + 1

    def write_label(self, ws, row, label, unit=None, typ=None, indent=0, bold=False, src=None, note=None):
        c = ws.cell(row, LABEL_COL)
        c.value = (" " + label) if isinstance(label, str) and label.startswith("=") else label
        c.font = S.F_BOLD if bold else S.F_BODY
        c.alignment = S.A_INDENT if indent == 1 else (S.A_INDENT2 if indent == 2 else S.A_LEFT)
        if unit:
            u = ws.cell(row, UNIT_COL)
            u.value = unit
            u.font = S.F_NOTE
            u.alignment = S.A_CENTER
        if typ:
            t = ws.cell(row, TYPE_COL)
            t.value = typ
            t.font = S.F_NOTE
            t.alignment = S.A_CENTER
        if src:
            s = ws.cell(row, SRC_COL)
            s.value = src
            s.font = S.F_NOTE
        if note:
            n = ws.cell(row, NOTE_COL)
            n.value = note
            n.font = S.F_NOTE

    def write_annual_formula(self, sheetkey, row, template, ctx, fmt, f0=None, years=None,
                             font=None, total=False, bold=False, fill=None, border=None):
        """Write a formula template across the year columns."""
        ws = self.ws(sheetkey)
        yrs = years or YEARS
        for y in yrs:
            col = YEAR_COL[y]
            tpl = f0 if (y == YEARS[0] and f0 is not None) else template
            if tpl is None or tpl == "":
                continue
            cell = ws.cell(row, col)
            if isinstance(tpl, (int, float)):
                cell.value = tpl
            else:
                cell.value = self.resolve(tpl, col, ctx) if tpl.startswith("=") else tpl
            cell.number_format = S.NF[fmt]
            cell.font = font or (S.F_BOLD if bold else S.F_BODY)
            cell.alignment = S.A_RIGHT
            if fill:
                cell.fill = fill
            if border:
                cell.border = border
        if total:
            t = ws.cell(row, SCALAR_COL)
            t.value = f"=SUM({FCOL}{row}:{LCOL}{row})" if total is True else self.resolve(total, SCALAR_COL, ctx)
            t.number_format = S.NF[fmt]
            t.font = S.F_BOLD if bold else S.F_BODY
            t.alignment = S.A_RIGHT
            if fill:
                t.fill = fill

    def write_annual_values(self, sheetkey, row, values, fmt, font=None, fill=None, default=None):
        ws = self.ws(sheetkey)
        for y in YEARS:
            col = YEAR_COL[y]
            v = values.get(y, default) if isinstance(values, dict) else values
            if v is None:
                continue
            cell = ws.cell(row, col)
            cell.value = v
            cell.number_format = S.NF[fmt]
            cell.font = font or S.F_INPUT
            cell.alignment = S.A_RIGHT
            if fill:
                cell.fill = fill

    def write_scalar(self, sheetkey, row, value, fmt, ctx=None, font=None, fill=None):
        ws = self.ws(sheetkey)
        cell = ws.cell(row, SCALAR_COL)
        if isinstance(value, str) and value.startswith("="):
            cell.value = self.resolve(value, SCALAR_COL, ctx) if ctx else value
        else:
            cell.value = value
        cell.number_format = S.NF[fmt]
        cell.font = font or S.F_INPUT
        cell.alignment = S.A_RIGHT
        if fill:
            cell.fill = fill
        return cell

    def shade_pre_valuation(self, sheetkey, row, upto_year=2026):
        ws = self.ws(sheetkey)
        for y in YEARS:
            if y <= upto_year:
                ws.cell(row, YEAR_COL[y]).fill = S.FILL_PRE

    def set_print(self, sheetkey, landscape=True, fit_width=True):
        ws = self.ws(sheetkey)
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.gridLines = False


CLASS_LABEL = {
    "v29": "v29 source",
    "DERIVED": "v29 derived",
    "UPDATED": "Updated info (user)",
    "CHANGE": "Modelling change (user request)",
    "DUMMY": "DUMMY — REPLACE WITH FS DATA",
    "PROV": "Provisional — to confirm",
    "CONTROL": "Control / switch",
    "CALC": "Formula",
}

CLASS_FILL = {
    "v29": S.FILL_INPUT,
    "DERIVED": S.FILL_INPUT,
    "UPDATED": S.FILL_UPDATED,
    "CHANGE": S.FILL_CHANGE,
    "DUMMY": S.FILL_DUMMY,
    "PROV": S.FILL_PROV,
    "CONTROL": S.FILL_CONTROL,
}
