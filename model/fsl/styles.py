"""Shared visual conventions for the FSL model workbook."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.formatting.rule import CellIsRule, FormulaRule

FONT_NAME = "Arial"
SZ = 9

def font(bold=False, color="000000", size=SZ, italic=False, underline=None):
    return Font(name=FONT_NAME, size=size, bold=bold, color=color, italic=italic, underline=underline)

F_BODY = font()
F_BOLD = font(bold=True)
F_TITLE = font(bold=True, size=14, color="1F3864")
F_SUB = font(italic=True, color="595959")
F_HDR = font(bold=True, color="FFFFFF")
F_INPUT = font(color="0000FF")                 # hard-coded input
F_INPUT_B = font(color="0000FF", bold=True)
F_LINK = font(color="00703C")                  # link from another sheet
F_LINK_B = font(color="00703C", bold=True)
F_NAV = font(color="0563C1", underline="single")
F_NOTE = font(italic=True, color="7F7F7F", size=8)
F_WARN = font(bold=True, color="C00000")
F_DUMMY = font(color="0000FF", bold=True)

FILL_HDR = PatternFill("solid", fgColor="1F3864")
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_SUBSECTION = PatternFill("solid", fgColor="F2F2F2")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")      # editable input
FILL_CONTROL = PatternFill("solid", fgColor="FFE699")    # switches / controls
FILL_DUMMY = PatternFill("solid", fgColor="F8CBAD")      # DUMMY - replace with FS data
FILL_CHANGE = PatternFill("solid", fgColor="DDEBF7")     # user-requested modelling change
FILL_UPDATED = PatternFill("solid", fgColor="E2EFDA")    # updated information supplied by user
FILL_PROV = PatternFill("solid", fgColor="FCE4D6")       # provisional / to confirm
FILL_TOTAL = PatternFill("solid", fgColor="EDEDED")
FILL_PASS = PatternFill("solid", fgColor="C6EFCE")
FILL_FAIL = PatternFill("solid", fgColor="FFC7CE")
FILL_WHITE = PatternFill("solid", fgColor="FFFFFF")
FILL_PRE = PatternFill("solid", fgColor="EDEDED")        # pre-valuation (historical / estimate) columns
FILL_KEY = PatternFill("solid", fgColor="FFFFCC")

THIN = Side(style="thin", color="BFBFBF")
MED = Side(style="medium", color="1F3864")
B_BOTTOM = Border(bottom=THIN)
B_TOP = Border(top=THIN)
B_TOTAL = Border(top=THIN, bottom=Side(style="double", color="1F3864"))
B_BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

A_LEFT = Alignment(horizontal="left", vertical="center")
A_RIGHT = Alignment(horizontal="right", vertical="center")
A_CENTER = Alignment(horizontal="center", vertical="center")
A_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
A_INDENT = Alignment(horizontal="left", vertical="center", indent=1)
A_INDENT2 = Alignment(horizontal="left", vertical="center", indent=2)

# number formats (zero shows as "-", negatives in parentheses)
NF = {
    "idr_bn": '#,##0.0;(#,##0.0);"-"',
    "idr_bn2": '#,##0.00;(#,##0.00);"-"',
    "usd_m": '#,##0.00;(#,##0.00);"-"',
    "usd_m1": '#,##0.0;(#,##0.0);"-"',
    "idr_m": '#,##0;(#,##0);"-"',
    "num": '#,##0;(#,##0);"-"',
    "num1": '#,##0.0;(#,##0.0);"-"',
    "num2": '#,##0.00;(#,##0.00);"-"',
    "num4": '#,##0.0000;(#,##0.0000);"-"',
    "tons": '#,##0;(#,##0);"-"',
    "idr_t": '#,##0;(#,##0);"-"',
    "idr": '#,##0;(#,##0);"-"',
    "pct": '0.0%;(0.0%);"-"',
    "pct2": '0.00%;(0.00%);"-"',
    "pct0": '0%;(0%);"-"',
    "pct4": '0.0000%;(0.0000%);"-"',
    "pct6": '0.000000%;(0.000000%);"-"',
    "x": '0.00"x";(0.00"x");"-"',
    "idx": '0.0000;(0.0000);"-"',
    "flag": '0;(0);"-"',
    "flag01": '0',
    "year": '0',
    "int": '0',
    "date": 'dd-mmm-yyyy',
    "text": '@',
    "days": '0',
    "fx": '#,##0',
    "sqm": '#,##0;(#,##0);"-"',
    "persons": '#,##0;(#,##0);"-"',
}

UNIT_FMT = {
    "IDR bn": "idr_bn", "USD m": "usd_m", "t": "tons", "tons": "tons", "IDR/t": "idr_t", "IDR/t/month": "idr_t",
    "%": "pct", "flag": "flag", "year": "year", "x": "x", "index": "idx", "days": "days", "date": "date",
    "IDR/USD": "fx", "sqm": "sqm", "IDR/sqm/month": "idr_t", "IDR m": "idr_m", "persons": "persons",
    "IDR m p.a.": "idr_m", "IDR bn p.a.": "idr_bn2", "months": "num1", "years": "int", "text": "text",
    "IDR": "idr", "USD": "num", "t/month": "tons", "0/1": "flag01", "1/2": "flag01",
}

def fmt_for(unit, fmt=None):
    if fmt:
        return NF[fmt]
    return NF[UNIT_FMT.get(unit, "num2")]
