"""Global Inputs sheet: timeline, valuation date, FX, tax, working capital, switches, valuation
parameters, project register, opening balance sheet, financing and dividend policy.

Row keys are registered on sheetkey 'global' and referenced from other sheets with {g:key}.
"""
import datetime as dt
from openpyxl.styles import Alignment
from . import styles as S
from .engine import (Model, YEARS, YEAR_COL, FIRST_COL, LAST_COL, FCOL, LCOL, L, yc, DATA_ROW0, SCALAR_COL,
                     SRC_COL, NOTE_COL, LABEL_COL, UNIT_COL, TYPE_COL, ROW_YEAR, GLOBAL_TITLE, CLASS_FILL)

TAB = "FFC000"

# (key, label, unit, kind, value(s), source, cls, note, fmt)
GENERAL = [
    ("val_date", "Valuation date", "date", "scalar", dt.date(2026, 12, 31), "User instruction (change from v29 which valued at model start)", "CHANGE",
     "Company DCF uses FCFF from the year after this date (FY2027+). Balances at this date enter the equity bridge.", "date"),
    ("val_year", "Valuation year (financial year ending on the valuation date)", "year", "calc", "=YEAR({g:val_date})", "Formula", "CALC", None, "year"),
    ("first_dcf_year", "First year of DCF cash flows", "year", "calc", "={g:val_year}+1", "Formula", "CALC", None, "year"),
    ("last_year", "Last explicit forecast year", "year", "calc", f"=MAX({FCOL}{ROW_YEAR}:{LCOL}{ROW_YEAR})", "Formula", "CALC", None, "year"),
    ("year", "Model year (header)", "year", "annual", {y: f"={yc(y)}${ROW_YEAR}" for y in YEARS}, "Timeline", "CALC", None, "year"),
    ("dcf_flag", "DCF window flag (1 = year after the valuation date)", "flag", "annual", {y: "=IF({y}>{g:val_year},1,0)" for y in YEARS}, "Formula", "CALC", None, "flag"),
    ("disc_period", "Discount period (years from the valuation date; year-end convention)", "years", "annual", {y: "=IF({y}>{g:val_year},{y}-{g:val_year},0)" for y in YEARS}, "Formula", "CALC",
     "FY2027 cash flow is discounted one full year from 31-Dec-2026.", "num"),
    ("df", "Discount factor at WACC (0 before / at the valuation date)", "x", "annual", {y: "=IF({y}>{g:val_year},1/(1+{g:wacc})^{g:disc_period},0)" for y in YEARS}, "Formula", "CALC", None, "num4"),
    ("pre_val_flag", "Pre-valuation flag (1 = up to and including the valuation year)", "flag", "annual", {y: "=IF({y}<={g:val_year},1,0)" for y in YEARS}, "Formula", "CALC", None, "flag"),
]

MACRO = [
    ("fx", "Exchange rate USD/IDR (applied to capex, USD presentation)", "IDR/USD", "annual", {y: 16265 for y in YEARS}, "Calculation_FSL!F17 (USD_IDR = 16,265, single rate in v29)", "v29",
     "v29 uses one rate for all years; each year is independently editable here.", "fx"),
    ("fx_val", "Exchange rate of the valuation year (USD presentation of values)", "IDR/USD", "calc", "=INDEX({g:fx@all},MATCH({g:val_year},{g:year@all},0))", "Formula", "CALC", None, "fx"),
    ("fx_open", "Exchange rate of the opening column (FY2023) — used for Dec-23 USD balances", "IDR/USD", "calc", "=INDEX({g:fx@all},1)", "Formula", "CALC", None, "fx"),
]

TAXWC = [
    ("tax_rate", "Corporate income tax rate", "%", "scalar", 0.22, "Output_Standalone (IDR)!J33 etc. (22% typed in each v29 tax formula)", "v29", "Applied per project line to MAX(0, EBIT) — v29 method, no loss relief between lines.", "pct"),
    ("ar_days", "Receivable days (on revenue)", "days", "scalar", 45, "Calculation_FSL!F1454", "v29", "Applied to every project line (v29 applies it to company revenue; result identical).", "days"),
    ("ap_days", "Payable days (on cost of sales)", "days", "scalar", 30, "Calculation_FSL!F1464", "v29", None, "days"),
    ("inventory", "Inventory balance (held at the Dec-23 level; centrally held, not allocated to projects)", "IDR bn", "scalar", 0.519811, "Output_FSL!I96 (BS_Historical!V10/1000)", "v29",
     "v29 holds this balance to FY2033 and drops it to nil from FY2034 through a formula break (Calculation_FSL!T1460:V1460); held flat here — see Reconciliation.", "idr_bn2"),
]

SWITCHES = [
    ("dep_method", "Depreciation method for project capex: 1 = v29 (total capex x share / life from commissioning, no end), 2 = by year of spend over the useful life from MAX(spend year, commissioning)", "1/2", "scalar", 2,
     "User request (capex % allocation linked to depreciation); v29 method retained as option 1", "CONTROL",
     "Method 2 links depreciation to the capex allocation and stops at the end of the useful life. Set 1 to reproduce v29 depreciation.", "flag01"),
    ("int_shield", "Corporate LT-loan interest tax shield at company level: 1 = deductible at the tax rate, 0 = no shield", "0/1", "scalar", 1,
     "Modelling judgement (v29 charged the interest to the Cigading existing / Wharf 2 lines and taxed each line separately)", "CONTROL", None, "flag01"),
    ("div_switch", "Dividend policy: 1 = v29 rules (30% of attributable profit, limited by reserve and cash before new equity), 0 = no dividends", "0/1", "scalar", 1,
     "Calculation_FSL rows 1527-1557", "CONTROL", None, "flag01"),
    ("funding_case", "Operating case underlying the Funding Requirement roll-forward: 1 = Without BUMN (conservative), 2 = With BUMN", "1/2", "scalar", 1,
     "User instruction: show which BUMN case underlies the funding calculation", "CONTROL", None, "flag01"),
    ("badas_in_funding", "Include Badas capex in the funding request while Badas is outside the valuation (1 = yes)", "0/1", "scalar", 1,
     "User instruction (separate funding-inclusion control)", "CONTROL", "When Badas is inside the valuation its capex is already in the company statements and this control is ignored.", "flag01"),
    ("other_bs_in_bridge", "Include Dec-23 other assets / other liabilities (held flat) in the equity bridge: 1 = yes, 0 = no", "0/1", "scalar", 0,
     "Modelling judgement — classification of these balances (deposits, deferred tax, accruals) not available", "PROV", None, "flag01"),
]

VALUATION = [
    ("rf", "IDR 10-year government bond yield", "%", "scalar", 0.0701, "Not in v29 — provisional market parameter", "PROV", None, "pct2"),
    ("sov_spread", "Indonesia sovereign default spread (removed from the bond yield)", "%", "scalar", 0.0162, "Not in v29 — provisional", "PROV", None, "pct2"),
    ("erp", "Equity risk premium incl. country risk", "%", "scalar", 0.0669, "Not in v29 — provisional", "PROV", None, "pct2"),
    ("beta_u", "Unlevered beta (port / logistics infrastructure)", "x", "scalar", 0.80, "Not in v29 — provisional", "PROV", None, "num2"),
    ("d_v", "Target debt / (debt + equity)", "%", "scalar", 0.20, "Not in v29 — provisional", "PROV", None, "pct"),
    ("kd_pre", "Pre-tax cost of debt", "%", "scalar", 0.105, "Call Rights!H71 note '10%-11%' (IDR) — provisional", "PROV", None, "pct2"),
    ("csrp", "Company-specific risk premium", "%", "scalar", 0.01, "Not in v29 — provisional", "PROV", None, "pct2"),
    ("rf_net", "Risk-free rate net of the sovereign spread", "%", "calc", "={g:rf}-{g:sov_spread}", "Formula", "CALC", None, "pct2"),
    ("d_e", "Target debt / equity", "x", "calc", "={g:d_v}/(1-{g:d_v})", "Formula", "CALC", None, "num2"),
    ("beta_l", "Relevered beta (Hamada)", "x", "calc", "={g:beta_u}*(1+(1-{g:tax_rate})*{g:d_e})", "Formula", "CALC", None, "num2"),
    ("ke", "Cost of equity = rf + beta x ERP + specific premium", "%", "calc", "={g:rf_net}+{g:beta_l}*{g:erp}+{g:csrp}", "Formula", "CALC", None, "pct2"),
    ("kd", "After-tax cost of debt", "%", "calc", "={g:kd_pre}*(1-{g:tax_rate})", "Formula", "CALC", None, "pct2"),
    ("wacc", "WACC (nominal IDR)", "%", "calc", "=(1-{g:d_v})*{g:ke}+{g:d_v}*{g:kd}", "Formula", "CALC", "Used by both valuation sheets and all project NPVs.", "pct2"),
    ("tg", "Terminal growth rate after the last forecast year (nominal IDR)", "%", "scalar", 0.03, "Not in v29 — provisional", "PROV", None, "pct2"),
    ("tv_capex_factor", "Terminal-year normalised capex as a multiple of terminal-year depreciation (1.00 = capex replaces depreciation)", "x", "scalar", 1.0, "Not in v29 — provisional", "PROV",
     "v29 forecasts no capex after FY2028; the terminal cash flow is normalised so that reinvestment sustains the asset base.", "num2"),
]

OPENING = [
    ("open_cash", "Cash and equivalents", "IDR bn", "scalar", 60.6389, "Output_FSL!I88 (BS_Historical!V7/1000)", "v29", None, "idr_bn"),
    ("open_inventory", "Inventory", "IDR bn", "scalar", "={g:inventory}", "Output_FSL!I96", "CALC", None, "idr_bn2"),
    ("open_ar", "Trade receivables (typed in v29; BS_Historical shows 104.1)", "IDR bn", "scalar", 45.0, "Output_FSL!I97 (typed value)", "v29", "Kept as v29; source issue — see Checks / Reconciliation.", "idr_bn"),
    ("open_fa", "Fixed assets, net (existing businesses; not allocated by business)", "IDR bn", "scalar", 826.022, "Output_FSL!I98 (BS_Historical!V17/1000)", "v29", None, "idr_bn"),
    ("open_other_assets_v29", "Other assets as reported in v29 (includes assets under construction of Wharf 2 / conveyor)", "IDR bn", "scalar", 331.263, "Output_FSL!I99 (BS_Historical!V23/1000)", "v29", None, "idr_bn"),
    ("open_ap", "Trade payables (typed in v29; BS_Historical shows 23.3)", "IDR bn", "scalar", 14.671985982904882, "Output_FSL!I103 (typed value)", "v29", None, "idr_bn"),
    ("open_loans", "Bank loans (LT loan 12 USD m + other loans)", "IDR bn", "scalar", 224.518, "Output_FSL!I104 (BS_Historical!V27/1000)", "v29", None, "idr_bn"),
    ("open_other_liab", "Other liabilities (held flat)", "IDR bn", "scalar", 55.9245, "Output_FSL!I105 (= BS_Historical!V28/1000 - 65.83 + 15.32, a v29 plug)", "v29", "Source issue — see Checks.", "idr_bn"),
    ("open_equity", "Paid-in capital", "IDR bn", "scalar", 677.731, "Output_FSL!I109 (BS_Historical!V32/1000)", "v29", None, "idr_bn"),
    ("open_re", "Retained earnings", "IDR bn", "scalar", 290.592, "Output_FSL!I111 (BS_Historical!V33/1000)", "v29", None, "idr_bn"),
    ("open_resid", "Opening balance residual = assets - liabilities - equity (v29 source imbalance, not plugged)", "IDR bn", "calc",
     "={g:open_cash}+{g:open_inventory}+{g:open_ar}+{g:open_fa}+{g:open_other_assets_v29}-{g:open_ap}-{g:open_loans}-{g:open_other_liab}-{g:open_equity}-{g:open_re}",
     "Formula", "CALC", "v29 hides this 0.006 residual with ROUNDDOWN; shown explicitly here and tolerated in the balance check.", "num4"),
]

FINANCING = [
    ("lt_loan_usd", "Existing long-term loan — principal (USD m); v29 'FSL Loan Repayment' 12 USD m, shown as repaid in FY2025 but NOT repaid", "USD m", "scalar", 12, "Input!K1237 / Calculation_FSL!K1511", "v29", None, "usd_m"),
    ("lt_loan_idr", "Existing long-term loan — principal (IDR bn) at the opening exchange rate", "IDR bn", "calc", "={g:lt_loan_usd}*{g:fx_open}/1000", "Calculation_FSL!K1513 (= 195.18)", "CALC", None, "idr_bn"),
    ("other_loans_idr", "Other bank loans at Dec-23 = total loans - LT loan; repaid in FY2024 as in v29", "IDR bn", "calc", "={g:open_loans}-{g:lt_loan_idr}", "Calculation_FSL!J1516 (29.34 repaid FY2024)", "CALC", None, "idr_bn"),
    ("lt_rate", "LT loan interest rate (p.a.) — implied by v29 FY2025 interest 15.27 / 195.18", "%", "scalar", 0.07825321, "Derived: (Input!L388 + Input!L433) / Calculation_FSL!K1513", "PROV",
     "Loan terms (rate, currency, maturity) to be confirmed; v29 shows no maturity.", "pct2"),
    ("fin_cost_2024", "FY2024 interest on all existing loans (v29 figure kept for the historical year)", "IDR bn", "scalar", 5.470905517661, "Input!K388 -> Output_Standalone (IDR)!J125", "v29", None, "idr_bn"),
    ("inj_date", "Capital injection — date proceeds are received (PROVISIONAL — confirm with the transaction timetable)", "date", "scalar", dt.date(2027, 6, 30),
     "User instruction: loan repaid when the injection is received; timing provisional", "PROV",
     "Moving this date moves the loan repayment, the interest charge and the year in which equity is received.", "date"),
    ("inj_year", "Capital injection — financial year of receipt", "year", "calc", "=YEAR({g:inj_date})", "Formula", "CALC", None, "year"),
    ("inj_frac", "Fraction of the injection year elapsed before receipt (interest accrues on the repaid amount for this fraction)", "%", "calc",
     "=(({g:inj_date}-DATE({g:inj_year},1,1))+1)/(DATE({g:inj_year},12,31)-DATE({g:inj_year},1,1)+1)", "Formula", "CALC", None, "pct"),
    ("inj_mode", "Capital injection amount: 1 = equal to the computed equity requirement (Funding Requirement), 2 = fixed amount below", "1/2", "scalar", 1,
     "User instruction: editable amount; mode 1 self-sizes to the funding gap", "CONTROL", None, "flag01"),
    ("inj_amount_usd", "Capital injection — fixed amount received (USD m), used only in mode 2", "USD m", "scalar", 0, "User input", "CONTROL",
     "Set mode = 2 to test a fixed amount (e.g. 0 to test that the loan is not extinguished without proceeds).", "usd_m"),
    ("repay_share", "Planned repayment of the LT loan out of the injection proceeds (% of the balance outstanding at the injection date)", "%", "scalar", 1.0,
     "Board deck Dec-2024 use of proceeds: LT loan repayment 12 USD m", "UPDATED", None, "pct"),
    ("min_cash", "Minimum operating cash balance", "IDR bn", "scalar", 3.0, "Calculation_FSL!J1531 (-3.00, sign reversed)", "v29", None, "idr_bn"),
    ("div_payout", "Dividend payout ratio (x profit attributable to FSL shareholders)", "%", "scalar", 0.30, "Calculation_FSL!F1550", "v29", None, "pct"),
    ("reserve_capital_usd", "Capital base for the statutory reserve test (USD m)", "USD m", "scalar", 18, "Calculation_FSL!F1537 (18 USD m)", "v29", None, "usd_m"),
    ("reserve_pct", "Minimum reserve as % of the capital base", "%", "scalar", 0.20, "Calculation_FSL!F1538", "v29", None, "pct"),
]



SECTIONS = [
    ("1. Timeline and valuation date", GENERAL),
    ("2. Exchange rate", MACRO),
    ("3. Tax and working capital", TAXWC),
    ("4. Model switches", SWITCHES),
    ("5. Valuation parameters (discount rate build-up, terminal value) — PROVISIONAL, not in v29", VALUATION),
    ("6. Opening balance sheet at 31-Dec-2023 (v29 Output_FSL column I) — company level; not allocated to projects", OPENING),
    ("7. Financing, capital injection, cash and dividend policy", FINANCING),
]


def plan_globals(model: Model, specs):
    """Phase 1: create the sheet and register every row (no formulas written yet)."""
    model.add_sheet("global", GLOBAL_TITLE, TAB)
    layout = []
    r = DATA_ROW0
    for title, rows in SECTIONS:
        layout.append((r, "section", title))
        r += 1
        for key, label, unit, kind, val, src, cls, note, fmt in rows:
            model.register("global", key, r, kind=("scalar" if kind in ("scalar", "calc") else "annual"), fmt=fmt, label=label)
            layout.append((r, "row", (key, label, unit, kind, val, src, cls, note, fmt)))
            r += 1
        r += 1
    # project register
    layout.append((r, "section", "8. Project register — inclusion in valuation, case membership, FSL economic interest (edit the blue cells)"))
    r += 1
    layout.append((r, "reg_header", None))
    r += 1
    reg_first = r
    for sp in specs:
        pid = sp["id"]
        model.register("global", f"incl.{pid}", r, kind="scalar", fmt="flag01", col=5)
        model.register("global", f"case.{pid}", r, kind="scalar", fmt="flag01", col=6)
        model.register("global", f"eint.{pid}", r, kind="scalar", fmt="pct", col=7)
        model.register("global", f"flagwo.{pid}", r, kind="scalar", fmt="flag01", col=11)
        model.register("global", f"flagw.{pid}", r, kind="scalar", fmt="flag01", col=12)
        layout.append((r, "reg_row", sp))
        r += 1
    model.register("global", "reg_first", reg_first, kind="scalar")
    model.register("global", "reg_last", r - 1, kind="scalar")
    layout.append((r, "reg_note", None))
    r += 2
    layout.append((r, "section", "9. Allocation bases (formulas) — Dec-23 company receivables / payables allocated to existing businesses pro rata FY2024 activity; assets under construction reclassified"))
    r += 1
    for key in ("rev2024_existing", "cogs2024_existing", "pre2024_capex_idr", "open_other_assets"):
        model.register("global", key, r, kind="scalar", fmt="idr_bn")
        layout.append((r, "alloc", key))
        r += 1
    r += 1
    layout.append((r, "legend", None))
    model._global_layout = layout
    return r


def write_globals(model: Model, specs):
    """Phase 2: write labels, values and formulas."""
    ws = model.ws("global")
    ws.column_dimensions["B"].width = 78
    model.write_ts_header("global", "GLOBAL INPUTS — shared assumptions, controls and project register",
                          "Single source for every assumption used by more than one sheet. Blue = input (yellow fill); bold blue on gold = control switch; orange = provisional; light blue = user-requested change. Black = formula.")
    ctx = {"sheet": "global", "pid": None}
    existing = [sp for sp in specs if sp.get("existing")]
    for r, typ, payload in model._global_layout:
        if typ == "section":
            model.section("global", r, payload)
        elif typ == "row":
            key, label, unit, kind, val, src, cls, note, fmt = payload
            model.write_label(ws, r, label, unit=unit, typ=cls if cls != "CALC" else "formula", src=src, note=note)
            ws.cell(r, TYPE_COL).font = S.font(bold=cls in ("CHANGE", "UPDATED", "PROV", "CONTROL"), color="595959", size=8)
            fill = CLASS_FILL.get(cls)
            if kind == "scalar":
                model.write_scalar("global", r, val, fmt, ctx=ctx, font=S.F_INPUT_B if cls == "CONTROL" else S.F_INPUT, fill=fill)
            elif kind == "calc":
                model.write_scalar("global", r, val, fmt, ctx=ctx, font=S.F_BODY)
            else:
                if isinstance(val, dict) and all(isinstance(v, str) and str(v).startswith("=") for v in val.values()):
                    if len(set(val.values())) == 1:
                        model.write_annual_formula("global", r, next(iter(val.values())), ctx, fmt, font=S.F_BODY)
                    else:
                        for y, f in val.items():
                            c = ws.cell(r, YEAR_COL[y])
                            c.value = model.resolve(f, YEAR_COL[y], ctx)
                            c.number_format = S.NF[fmt]
                            c.font = S.F_BODY
                            c.alignment = S.A_RIGHT
                else:
                    model.write_annual_values("global", r, val, fmt, fill=fill)
        elif typ == "reg_header":
            hdr = ["ID  Project", "Entity", "Category", "In valuation\n(1/0)", "Case\n(1 both / 2 With-BUMN only)", "FSL econ.\ninterest", "Commissioning\nyear",
                   "Total capex\nUSD m", "Capex after\nval. date USD m", "Included\nw/o BUMN", "Included\nw/ BUMN", "Input", "Calc", "Output"]
            for i, h in enumerate(hdr):
                c = ws.cell(r, 2 + i)
                c.value = h
                c.font = S.F_HDR
                c.fill = S.FILL_HDR
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[r].height = 40
        elif typ == "reg_row":
            sp = payload
            pid = sp["id"]
            d = sp.get("defaults", {})
            ws.cell(r, 2).value = f"{pid}  {sp['name']}"
            ws.cell(r, 2).font = S.F_BODY
            ws.cell(r, 3).value = sp.get("entity_short", sp.get("entity", ""))
            ws.cell(r, 3).font = S.F_NOTE
            ws.cell(r, 4).value = sp.get("category", "")
            ws.cell(r, 4).font = S.F_NOTE
            for col, val, fmt in [(5, d.get("in_valuation", 1), "flag01"), (6, d.get("case", 1), "flag01"), (7, d.get("econ_interest", 1.0), "pct")]:
                c = ws.cell(r, col)
                c.value = val
                c.number_format = S.NF[fmt]
                c.font = S.F_INPUT_B
                c.fill = S.FILL_CONTROL
                c.alignment = S.A_CENTER
            tin, tcalc, tout = f"{pid} {sp['code']} Input", f"{pid} {sp['code']} Calc", f"{pid} {sp['code']} Output"
            c = ws.cell(r, 8)
            if sp.get("commissioning"):
                c.value = "=" + model.cell(f"{pid}:in", "commissioning", cur_sheet="global")
            else:
                c.value = "existing"
            c.font = S.F_LINK
            c.alignment = S.A_CENTER
            c.number_format = S.NF["year"]
            for col, key in [(9, "capex_total_usd"), (10, "capex_post_val")]:
                c = ws.cell(r, col)
                c.value = ("=" + model.cell(f"{pid}:in", key, cur_sheet="global")) if sp.get("capex") else 0
                c.font = S.F_LINK
                c.number_format = S.NF["usd_m"]
            c = ws.cell(r, 11)
            c.value = f"={L(5)}{r}*IF({L(6)}{r}=1,1,0)"
            c.number_format = S.NF["flag01"]
            c.alignment = S.A_CENTER
            c = ws.cell(r, 12)
            c.value = f"={L(5)}{r}"
            c.number_format = S.NF["flag01"]
            c.alignment = S.A_CENTER
            for col, t in [(13, tin), (14, tcalc), (15, tout)]:
                c = ws.cell(r, col)
                c.value = "open"
                c.hyperlink = f"#'{t}'!A1"
                c.font = S.F_NAV
                c.alignment = S.A_CENTER
        elif typ == "reg_note":
            ws.cell(r, 2).value = ("Case 1 = the project is in both company cases; case 2 = only in the With-BUMN case (BUMN operatorship schemes). "
                                   "'In valuation' = 0 removes the project from both cases, the funding roll-forward and the valuations (its sheets keep calculating for information).")
            ws.cell(r, 2).font = S.F_NOTE
        elif typ == "alloc":
            key = payload
            rev_terms = "+".join(f"{{x:{sp['id']}:{sp['hooks']['revenue']}@2024}}" for sp in existing) or "1"
            cogs_terms = "+".join(f"{{x:{sp['id']}:{sp['hooks']['cogs']}@2024}}" for sp in existing if sp["hooks"].get("cogs")) or "1"
            capex_terms = "+".join(f"{{x:{sp['id']}:capex_idr@2023}}" for sp in specs if sp.get("capex")) or "0"
            labels = {"rev2024_existing": ("FY2024 revenue of existing businesses (sum)", "=" + rev_terms),
                      "cogs2024_existing": ("FY2024 cost of sales of existing businesses (sum)", "=" + cogs_terms),
                      "pre2024_capex_idr": ("Pre-2024 project capex (assets under construction at Dec-23) reclassified from other assets to project fixed assets", "=" + capex_terms),
                      "open_other_assets": ("Other assets at Dec-23 excluding the reclassified assets under construction", "={g:open_other_assets_v29}-{g:pre2024_capex_idr}")}
            label, f = labels[key]
            model.write_label(ws, r, label, unit="IDR bn", typ="formula")
            model.write_scalar("global", r, f, "idr_bn", ctx=ctx, font=S.F_LINK)
        elif typ == "legend":
            model.section("global", r, "Legend — input classification (column D)")
            r2 = r + 1
            legend = [("v29", "v29 source — value and cell reference from V29_Ruby_FM_Danantara.xlsm (effective value used by the v29 calculation)"),
                      ("DERIVED", "v29 derived — computed from v29 figures (e.g. implied interest rate)"),
                      ("UPDATED", "Updated information supplied by the user (e.g. board-deck use of proceeds)"),
                      ("CHANGE", "User-requested modelling change (valuation date, injection-linked repayment, capex %, corrections of v29 formula breaks)"),
                      ("DUMMY", "Dummy assumption — REPLACE WITH FS DATA (Badas working paper)"),
                      ("PROV", "Provisional — to confirm (timing, rates, market parameters)"),
                      ("CONTROL", "Control / switch")]
            for cls, text in legend:
                ws.cell(r2, 4).value = cls
                ws.cell(r2, 4).fill = CLASS_FILL.get(cls, S.FILL_WHITE)
                ws.cell(r2, 4).font = S.font(size=8, bold=True)
                ws.cell(r2, 4).alignment = S.A_CENTER
                ws.cell(r2, 5).value = text
                ws.cell(r2, 5).font = S.F_NOTE
                r2 += 1
    model.set_print("global")
