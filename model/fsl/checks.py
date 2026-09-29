"""Checks sheet (formula-based PASS/FAIL) and Reconciliation to v29 (static v29 reference values vs the revised With-BUMN case)."""
import os, warnings
import openpyxl
from . import styles as S
from .engine import Model, YEARS, YEAR_COL, FCOL, LCOL, L, SCALAR_COL, ROW_YEAR, DATA_ROW0, LABEL_COL
from .layout import plan_layout, write_layout

T_CHECKS = "Checks"
T_RECON = "Reconciliation v29"
TAB = "C00000"
V29_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "source", "V29_Ruby_FM_Danantara.xlsm")
V29_COL = {y: 10 + (y - 2024) for y in range(2024, 2037)}   # Output_FSL J..V


def _check(key, label, formula, tol, kind="check", note=None):
    """A check row: value formula in E, tolerance in F, status in G."""
    return ("custom", (lambda key, label, formula, tol, kind, note: (lambda m, sk, r, ctx: _write_check(m, sk, r, ctx, key, label, formula, tol, kind, note)))(key, label, formula, tol, kind, note))


def _write_check(m, sk, r, ctx, key, label, formula, tol, kind, note):
    ws = m.ws(sk)
    ws.cell(r, 2).value = label
    ws.cell(r, 2).font = S.F_BODY
    ws.cell(r, 2).alignment = S.A_WRAP
    c = ws.cell(r, 5)
    c.value = m.resolve(formula, 5, ctx)
    c.number_format = S.NF["num4"]
    c.font = S.F_BODY
    c.alignment = S.A_RIGHT
    t = ws.cell(r, 6)
    t.value = tol
    t.number_format = S.NF["num4"]
    t.font = S.F_INPUT
    t.alignment = S.A_RIGHT
    s = ws.cell(r, 7)
    if kind == "check":
        s.value = f'=IF(ABS(E{r})<=F{r},"PASS","FAIL")'
    elif kind == "warn":
        s.value = f'=IF(ABS(E{r})<=F{r},"OK","WARNING")'
    else:
        s.value = f'=IF(ABS(E{r})<=F{r},"INFO","INFO")'
    s.font = S.F_BOLD
    s.alignment = S.A_CENTER
    ws.cell(r, 8).value = kind.upper()
    ws.cell(r, 8).font = S.F_NOTE
    if note:
        ws.cell(r, 9).value = note
        ws.cell(r, 9).font = S.F_NOTE
    m.register(sk, key, r, kind="scalar", fmt="text", col=7)


def _fs_span(fs, key, y0=2024, y1=2036):
    return f"{{s:{fs}:{key}@all}}"


class ChecksPlan:
    def __init__(self, model: Model, specs):
        self.m = model
        self.specs = specs
        model.add_sheet("checks", T_CHECKS, TAB, ts=False, widths={"A": 2, "B": 95, "C": 14, "D": 14, "E": 14, "F": 12, "G": 12, "H": 12, "I": 70, "J": 12, "K": 12, "L": 12})
        model.add_sheet("recon", T_RECON, TAB)
        rows = [
            ("custom", lambda m, sk, r, ctx: _overall(m, sk, r, ctx)),
            ("blank",),
            ("custom", lambda m, sk, r, ctx: _hdr(m, sk, r)),
            _check("c_bs_wo", "Balance sheet balances — Without BUMN (max |assets - liabilities - equity|, FY2024-36)", "=SUMPRODUCT(MAX(ABS({s:fs_wo:bs_check@fc})))", 0.01, note="2023 opening residual 0.006 is a v29 source imbalance (typed receivables / payables)"),
            _check("c_bs_w", "Balance sheet balances — With BUMN", "=SUMPRODUCT(MAX(ABS({s:fs_w:bs_check@fc})))", 0.01),
            _check("c_cash_wo", "Cash: closing cash = pre-equity path + cumulative injection — Without BUMN", "=SUMPRODUCT(MAX(ABS({s:fs_wo:cash_check@fc})))", 0.001),
            _check("c_cash_w", "Cash: closing cash = pre-equity path + cumulative injection — With BUMN", "=SUMPRODUCT(MAX(ABS({s:fs_w:cash_check@fc})))", 0.001),
            _check("c_rev_wo", "Consolidation: sum of project revenue x flags = company revenue — Without BUMN", "=SUMPRODUCT(MAX(ABS({s:summary:chk_rev_wo@all})))", 0.001),
            _check("c_rev_w", "Consolidation: sum of project revenue x flags = company revenue — With BUMN", "=SUMPRODUCT(MAX(ABS({s:summary:chk_rev_w@all})))", 0.001),
            _check("c_fcff_w", "Consolidation: sum of project FCFF x flags = valuation FCFF — With BUMN", "=SUMPRODUCT(MAX(ABS({s:summary:chk_fcff_w@all})))", 0.001),
            _check("c_capex_pct", "Capex allocation: number of projects whose spending % do not add to 100%",
                   "=" + "+".join(f'IF({{xin:{sp["id"]}:capex_check}}="PASS",0,1)' for sp in specs if sp.get("capex")), 0),
            _check("c_capex_unalloc", "Capex allocation: unallocated capex across projects (USD m)", "=" + "+".join(f"ABS({{xin:{sp['id']}:capex_unalloc}})" for sp in specs if sp.get("capex")), 0.0001),
            _check("c_val_date", "Valuation date = 31-Dec-2026 (days difference)", "={g:val_date}-DATE(2026,12,31)", 0, note="Change the date on Global Inputs deliberately; this check documents the delivered setting"),
            _check("c_dcf_start", "DCF window starts the year after the valuation date (sum of discount factors up to the valuation year)", "=SUMPRODUCT({g:df@all},--({g:year@all}<={g:val_year}))", 0),
            _check("c_df_2027", "First DCF year discounted one full year: DF(first year) - 1/(1+WACC)", "=INDEX({g:df@all},MATCH({g:first_dcf_year},{g:year@all},0))-1/(1+{g:wacc})", 0.000001),
            _check("c_lt_nonneg", "LT loan balance never negative (repayment <= balance)", "=MIN(0,MIN({s:fin:lt_close@all}))", 0),
            _check("c_lt_carried", "LT loan not extinguished without proceeds: if total injection = 0 the FY2036 balance equals the principal", "=IF(SUM({s:fin:inj_idr@all})=0,{s:fin:lt_close@2036}-{g:lt_loan_idr},0)", 0.001),
            _check("c_gap_mode1", "Funding: residual gap after the injection in the funding case (must be nil in mode 1)", "=IF({g:inj_mode}=1,{s:fund:resid_gap_sel},0)", 0.001),
            _check("c_inj_ge_repay", "Funding: in mode 1 the injection covers the planned LT loan repayment", "=IF({g:inj_mode}=1,MIN(0,SUM({s:fin:inj_idr@all})-SUM({s:fin:lt_planned@all})),0)", 0.001),
            _check("c_div_wo", "Dividends never negative — Without BUMN", "=MIN(0,-MAX({s:fs_wo:div_paid@all}))", 0),
            _check("c_div_w", "Dividends never negative — With BUMN", "=MIN(0,-MAX({s:fs_w:div_paid@all}))", 0),
            _check("c_open_bs", "Opening balance sheet residual (v29 source imbalance, tolerated)", "={g:open_resid}", 0.01, note="v29 Output_FSL hides a 0.006 residual with ROUNDDOWN; kept visible"),
            _check("c_pv_wo", "Valuation: sum of project explicit PVs = company explicit PV — Without BUMN", "={s:val_wo:pv_check}", 0.001),
            _check("c_pv_w", "Valuation: sum of project explicit PVs = company explicit PV — With BUMN", "={s:val_w:pv_check}", 0.001),
            _check("c_ph_ok", "Funding-only placeholders: active placeholders with % not adding to 100%", "=" + "+".join(f'IF({{s:fund:F0{i}_check}}="CHECK: % not 100%",1,0)' for i in range(1, 6)), 0),
            _check("c_badas_val", "Badas outside the headline valuation while dummy (register switch, delivered = 0)", "={g:incl.P19}" if any(sp["id"] == "P19" for sp in specs) else "=0", 0, kind="info"),
            _check("w_cash_pre", "Bridge need before the injection (largest pre-injection shortfall vs minimum cash, funding case)", "={s:fund:bridge_need}", 0.001, kind="warn",
                   note="v29 assumed equity in FY2026-27 matching capex; with the injection after 2026 the FY2026 capex needs bridge funding until the injection is received"),
            _check("w_nbv_neg", "Project fixed assets: most negative project NBV (method 1 can depreciate beyond cost — v29 behaviour)", "=MIN(0," + ",".join(f"MIN({{x:{sp['id']}:nbv_new@all}})" for sp in specs if sp.get("capex")) + ")", 0.001, kind="warn"),
            _check("w_fa_exist", "Existing fixed assets NBV at FY2036 (v29 depreciates existing assets beyond their Dec-23 book value: negative = source issue)", "=MIN(0,{s:fs_w:fa_exist@2036})", 0.001, kind="warn"),
            _check("w_lt_short", "LT loan repayment shortfall (mode 2 with insufficient proceeds)", "=SUM({s:fin:lt_shortfall@all})", 0.001, kind="warn"),
            _check("w_gap_other", "Residual funding gap in the non-funding case after the injection", "=IF({g:funding_case}=1,SUM({s:fs_w:gap_after@all}),SUM({s:fs_wo:gap_after@all}))", 0.001, kind="warn"),
            ("blank",),
            ("custom", lambda m, sk, r, ctx: _issues(m, sk, r)),
        ]
        self.planned, _ = plan_layout(model, "checks", rows, start=5)
        # reconciliation rows
        self.v29 = _load_v29()
        rrows = [("section", "A. Company income statement — v29 Output_FSL (With BUMN, Sapphire excluded) vs revised model With-BUMN case (IDR bn). v29 values are static references.")]
        for key, label, v29row, fskey, sign, expl in RECON_LINES:
            rrows += [
                ("values", f"v29_{key}", f"v29: {label}", "IDR bn", {y: self.v29["Output_FSL"][v29row].get(y, 0) for y in range(2024, 2037)}, "idr_bn", {"cls": "v29", "src": f"Output_FSL row {v29row} (static copy)"}),
                ("row", f"rev_{key}", f"Revised: {label}", "IDR bn", (f"={sign}{{s:fs_w:{fskey}}}" if fskey else "=0"), "idr_bn", {"link": True}),
                ("row", f"d_{key}", f"Difference (revised - v29)", "IDR bn", f"={{t:rev_{key}}}-{{t:v29_{key}}}", "idr_bn", {"bold": True, "note": expl}),
            ]
        rrows += [("blank",), ("section", "B. Project lines — v29 Output_Standalone (IDR) EBIT vs project Output EBIT (IDR bn); differences arise only from the depreciation method (switch = 1 reproduces v29) and documented v29 breaks")]
        for sp in specs:
            r = sp.get("v29_rows", {})
            pid = sp["id"]
            if r.get("ebit") and not sp.get("v29_calc_rows"):
                vals = self.v29["standalone"].get(r["ebit"], {})
                rrows += [
                    ("values", f"v29_ebit_{pid}", f"v29 EBIT — {pid} {sp['name']}", "IDR bn", {y: vals.get(y, 0) for y in range(2024, 2037)}, "idr_bn", {"cls": "v29", "src": f"Output_Standalone (IDR) row {r['ebit']} (static copy)"}),
                    ("row", f"ebit_{pid}", f"Revised EBIT — {pid}", "IDR bn", f"={{xo:{pid}:ebit}}", "idr_bn", {"link": True}),
                    ("row", f"d_ebit_{pid}", "Difference", "IDR bn", f"={{t:ebit_{pid}}}-{{t:v29_ebit_{pid}}}", "idr_bn", {"bold": True, "note": sp.get("recon_note", "Depreciation method 2 vs v29 method 1 (see Global Inputs switch)")}),
                ]
            elif sp.get("v29_calc_rows"):
                rrows.append(("note", f"{pid} {sp['name']}: v29 standalone block is zero (excluded by v29 flag) — see the project Output sheet; v29 Calculation_FSL values were matched in the build tests."))
            else:
                rrows.append(("note", f"{pid} {sp['name']}: not in v29 (new project template)."))
        rrows += [("blank",), ("section", "C. Explanation of differences (summary)"),
                  ("note", "1) Valuation date 31-Dec-2026: does not change the statements; only the DCF window and the balances used in the equity bridge."),
                  ("note", "2) LT loan: v29 repays 195.2 in FY2025; revised carries it to the injection date -> financing cost from FY2026, tax shield at company level, loan on the balance sheet until repaid, repayment funded by the injection."),
                  ("note", "3) Equity injection: v29 injects FY2026 (419.8) and FY2027 (160.0) = capex of those years ('Ass Oct 2025' scenario, Call Rights row 48); revised injects once, at the injection date, the requirement of the funding case."),
                  ("note", "4) Depreciation: default method 2 depreciates each year's capex over its life from MAX(spend year, commissioning) and stops at the end of life; v29 (method 1) depreciates total capex from commissioning without end. KBS warehouse: v29 depreciates nil through a broken reference; revised depreciates the 6.02 USD m."),
                  ("note", "5) Inventory: v29 drops the 0.52 balance to nil from FY2034 (formula break); revised holds it flat. Other assets: 198.4 of pre-2024 project spend reclassified from other assets to project fixed assets (total assets unchanged)."),
                  ("note", "6) Dividends: same rules (30% payout, reserve, cash cap) but the cash cap uses cash before the injection so dividends are not paid out of new equity; reserve test uses actual retained earnings."),
                  ("note", "7) Working capital: identical days; allocated to projects; company totals equal v29 by construction. Corporate interest tax shield: v29 taxed each block separately (shield only where the block was profitable)."),
                  ("note", "Set Global Inputs: depreciation method = 1, injection date = 31-Dec-2025, injection mode = 2 with 12 USD m to reproduce v29's income statement through EBIT and its FY2025 repayment (equity and cash then differ only by the injection basis).")]
        self.rplanned, _ = plan_layout(model, "recon", rrows)

    def write(self):
        m = self.m
        ws = m.ws("checks")
        ws["B1"] = "CHECKS — mechanical integrity (live formulas)"
        ws["B1"].font = S.F_TITLE
        ws["B2"] = "PASS / FAIL checks with tolerances (column F, editable). WARNING rows flag results that need attention but are not mechanical errors; INFO rows document settings."
        ws["B2"].font = S.F_SUB
        m.nav_links("checks", row=3)
        ctx = {"sheet": "checks", "pid": None}
        write_layout(m, "checks", self.planned, ctx)
        m.set_print("checks", landscape=False)
        m.write_ts_header("recon", "RECONCILIATION TO v29 — With-BUMN case vs V29 Output_FSL (Sapphire excluded in both)",
                          "v29 values are static copies (grey input style, source cited). Differences are explained in the note column of each difference row and in section C.")
        ctx = {"sheet": "recon", "pid": None}
        write_layout(m, "recon", self.rplanned, ctx)
        m.set_print("recon")


def _hdr(m, sk, r):
    ws = m.ws(sk)
    for col, h in [(2, "Check"), (5, "Value"), (6, "Tolerance"), (7, "Status"), (8, "Type"), (9, "Note")]:
        c = ws.cell(r, col)
        c.value = h
        c.font = S.F_HDR
        c.fill = S.FILL_HDR


def _overall(m, sk, r, ctx):
    ws = m.ws(sk)
    ws.cell(r, 2).value = "OVERALL STATUS (all PASS/FAIL checks)"
    ws.cell(r, 2).font = S.F_BOLD
    c = ws.cell(r, 7)
    c.value = f'=IF(COUNTIF(G{r+3}:G{r+60},"FAIL")=0,"ALL PASS","FAIL: "&COUNTIF(G{r+3}:G{r+60},"FAIL"))'
    c.font = S.F_BOLD
    c.alignment = S.A_CENTER
    m.register(sk, "overall", r, kind="scalar", fmt="text", col=7)
    c2 = ws.cell(r, 9)
    c2.value = f'=COUNTIF(G{r+3}:G{r+60},"WARNING")&" warning(s)"'
    c2.font = S.F_NOTE


ISSUES = [
    ("S1", "Opening balance sheet: v29 types receivables 45.0 and payables 14.672 (BS_Historical shows 104.1 / 23.3) and plugs other liabilities; residual 0.006 hidden by ROUNDDOWN. Kept as v29; shown in the Checks."),
    ("S2", "Existing-asset depreciation runs for the whole forecast without a net-book-value cap (existing fixed assets turn negative by FY2035 in v29). Kept; flagged as a warning."),
    ("S3", "KBS warehouse revitalisation: 6.02 USD m capex never depreciated in v29 (Calculation_FSL!F461 -> empty Input!L1270). Revised model depreciates it (difference explained in the Reconciliation)."),
    ("S4", "DMS trucking: equipment (0.04/0.04/0.03/0.04 USD m) depreciated in v29 with no capex cash flow. Reproduced as v29 and documented in the P08 sheets."),
    ("S5", "Inventory drops to nil from FY2034 in v29 (formula break); held flat in the revised model."),
    ("S6", "Belawan expansion volume ends in FY2034 (Calculation_SOEs row 450 empty in FY2035-36) — source gap kept; flagged in E3 inputs."),
    ("S7", "WIN Makassar sugar tariff FY2034-36 references an empty row (nil revenue) — kept; flagged in E4 inputs."),
    ("S8", "Teluk Lamong G&A staff cost grown with the legal-fee growth row (3% instead of 3.5%) — kept; flagged in E1 inputs."),
    ("S9", "Cigading existing volume: v29 overrides the input (1.99 m t) with a typed 1.1 m t; tariff from FY2025 linked to the Wharf 2 tariff — kept and cited."),
    ("S10", "LT loan terms (rate implied 7.83%, maturity, currency) and the capital injection date are provisional; FY2024-26 figures are v29 projections, not actuals."),
    ("S11", "Minority interest: v29 applies FSL's 65% NPLOG share only in the dividend base; the revised model shows the non-controlling share as a memo and deducts it in the equity bridge."),
]


def _issues(m, sk, r):
    ws = m.ws(sk)
    ws.cell(r, 2).value = "Genuine v29 source issues — documented, not concealed (see Reconciliation and the project Input notes)"
    ws.cell(r, 2).font = S.F_BOLD
    for i, (code, text) in enumerate(ISSUES):
        ws.cell(r + 1 + i, 2).value = f"{code}  {text}"
        ws.cell(r + 1 + i, 2).font = S.F_NOTE
        ws.cell(r + 1 + i, 2).alignment = S.A_WRAP
        ws.row_dimensions[r + 1 + i].height = 26


RECON_LINES = [
    # key, label, v29 Output_FSL row, fs key, sign, explanation
    ("rev", "Net sales", 44, "revenue", "", "Identical unless project inputs are changed (DMS / BUMN / all lines included in both)."),
    ("cogs", "Cost of sales", 45, "cogs", "", "Identical unless inputs change."),
    ("opex", "Operating expenses", 48, "opex", "", "Identical unless inputs change."),
    ("ebitda", "EBITDA", 49, "ebitda", "", "Identical unless inputs change."),
    ("da", "Depreciation & amortisation", 51, "da", "", "Depreciation method 2 (by year of spend, capped at useful life) vs v29 method 1; KBS warehouse depreciation added (v29 nil through a broken reference). Set method = 1 to compare."),
    ("ebit", "EBIT", 52, "ebit", "", "Follows D&A."),
    ("fin", "Financing cost", 54, "fin_cost", "", "LT loan carried until the injection date (v29: repaid FY2025). FY2024-25 identical (v29 figures / implied rate)."),
    ("tax", "Tax", 57, "tax", "", "Tax on EBIT per line identical to v29 method; corporate interest shield at company level (v29: shield only inside the Cigading / Wharf 2 lines)."),
    ("np", "Net profit", 58, "np", "", "Follows financing cost, tax and D&A."),
    ("capex", "Fixed asset acquisition (capex)", 77, "cf_capex", "", "Identical timing unless capex % are moved; v29 excludes pre-2024 spend from the cash flow, as here."),
    ("inj", "Equity injection", 81, "cf_inj", "", "v29: FY2026 + FY2027 capex-matching tranches; revised: single injection at the injection date sized to the funding requirement."),
    ("repay", "Loan repayment", 82, "cf_repay", "", "v29 repays the LT loan in FY2025; revised repays at the injection date."),
    ("div", "Dividend payment", 83, "cf_div", "", "Same policy; cash cap excludes the injection; reserve test on actual retained earnings."),
    ("cash", "Closing cash", 88, "cf_close", "", "Cumulative effect of the above."),
    ("fa", "Fixed assets (total)", 98, None, "", "Revised total = existing (unallocated) + project NBV incl. 198.4 reclassified assets under construction (v29 kept them in other assets)."),
    ("loan", "Bank loan", 104, "debt", "", "LT loan outstanding until the injection date."),
    ("equity", "Paid-in capital", 109, "paid_in", "", "Injection timing / amount."),
    ("re", "Retained earnings", 111, "re", "", "Follows net profit and dividends."),
]


def _load_v29():
    """Static v29 reference values (Output_FSL and Output_Standalone (IDR)) read once at build time."""
    out = {"Output_FSL": {}, "standalone": {}}
    if not os.path.exists(V29_PATH):
        return out
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(V29_PATH, data_only=True, read_only=True, keep_vba=True)
    ws = wb["Output_FSL"]
    rows = {r: vals for r, vals in zip(range(1, 140), ws.iter_rows(min_row=1, max_row=139, min_col=10, max_col=22, values_only=True))}
    for key, label, row, *_ in RECON_LINES:
        out["Output_FSL"][row] = {y: (rows[row][c - 10] or 0) for y, c in V29_COL.items()}
    ws2 = wb["Output_Standalone (IDR)"]
    for r, vals in zip(range(18, 527), ws2.iter_rows(min_row=18, max_row=526, min_col=10, max_col=22, values_only=True)):
        out["standalone"][r] = {y: (vals[c - 10] or 0) for y, c in V29_COL.items()}
    wb.close()
    return out
