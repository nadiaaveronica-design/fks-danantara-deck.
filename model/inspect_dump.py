#!/usr/bin/env python3
"""Dump every sheet of the source workbooks to TSV grids (values and formulas) under model/inspect/ for inspection.
Usage: python3 model/inspect_dump.py   (regenerates model/inspect/v29 and model/inspect/v8; not committed to git)"""
import openpyxl, os, re, warnings
warnings.filterwarnings("ignore")
from openpyxl.utils import get_column_letter
ROOT = os.path.dirname(os.path.abspath(__file__))


def safe(s):
    return re.sub(r'[^A-Za-z0-9_]+', '_', s)


def dump(path, outdir, vba):
    os.makedirs(outdir, exist_ok=True)
    wbf = openpyxl.load_workbook(path, keep_vba=vba)
    wbv = openpyxl.load_workbook(path, data_only=True, keep_vba=vba)
    for ws in wbf.worksheets:
        wv = wbv[ws.title]
        maxc = ws.max_column
        with open(f"{outdir}/{safe(ws.title)}__values.tsv", "w") as fv, open(f"{outdir}/{safe(ws.title)}__formulas.tsv", "w") as ff, open(f"{outdir}/{safe(ws.title)}__cells.txt", "w") as fc:
            hdr = "row\t" + "\t".join(get_column_letter(c) for c in range(1, maxc + 1)) + "\n"
            fv.write(hdr)
            ff.write(hdr)
            for r in range(1, ws.max_row + 1):
                vals, forms, any_ = [], [], False
                for c in range(1, maxc + 1):
                    f, v = ws.cell(r, c).value, wv.cell(r, c).value
                    if v is None and f is None:
                        vals.append("")
                        forms.append("")
                        continue
                    any_ = True
                    vs = v if v is not None else ""
                    if isinstance(vs, float):
                        vs = f"{vs:.6g}"
                    vals.append(str(vs).replace("\t", " ").replace("\n", " "))
                    fs = str(f).replace("\t", " ").replace("\n", " ") if f is not None else ""
                    forms.append(fs)
                    fc.write(f"{get_column_letter(c)}{r}\t{fs}\t{vs}\n")
                if any_:
                    fv.write(f"{r}\t" + "\t".join(vals) + "\n")
                    ff.write(f"{r}\t" + "\t".join(forms) + "\n")
    print("done", path)


if __name__ == "__main__":
    dump(os.path.join(ROOT, "source", "V29_Ruby_FM_Danantara.xlsm"), os.path.join(ROOT, "inspect", "v29"), True)
    dump(os.path.join(ROOT, "source", "FSL_Model_v8_Revised_2.xlsx"), os.path.join(ROOT, "inspect", "v8"), False)
