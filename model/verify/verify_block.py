#!/usr/bin/env python3
"""Verification harness: build a mini workbook with Global Inputs + one project (and the projects it
depends on), recalculate with LibreOffice, and compare the project's income-statement lines with the
v29 Output_Standalone (IDR) block, year by year (FY2024-FY2036).

Usage:  python3 model/verify/verify_block.py P02 [--dep-method 1] [--tol 0.001] [--keep]
Exit code 0 when every compared line matches within tolerance.
"""
import sys, os, json, argparse, importlib, subprocess, warnings
warnings.filterwarnings("ignore")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))     # model/
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.dirname(ROOT))
import openpyxl
from fsl.engine import Model, YEARS, YEAR_COL, L
from fsl.project import ProjectPlan, write_project
from fsl.globals_sheet import plan_globals, write_globals
from fsl.specs import load_spec, all_spec_ids

V29 = os.path.join(ROOT, "source", "V29_Ruby_FM_Danantara.xlsm")
RECALC = "/root/.claude/skills/synced/ca1271e6-ddce-4c4e-8f01-c832fbb88730_653470ac-0907-46f7-93c1-15505bb7874c/xlsx/scripts/recalc.py"
STANDALONE = "Output_Standalone (IDR)"
V29_YEAR_COL = {y: 10 + (y - 2024) for y in range(2024, 2037)}   # J=2024 .. V=2036


def resolve_deps(pid, seen=None):
    seen = [] if seen is None else seen
    if pid in seen:
        return seen
    seen.append(pid)
    sp = load_spec(pid)
    for d in sp.get("depends", []):
        resolve_deps(d, seen)
    return seen


def build(pids, path, dep_method=1, fx=16265):
    m = Model()
    specs = [load_spec(p) for p in pids]
    # globals first (needs specs for the register), then project plans, then writes
    # plan projects first so that register links can resolve
    plan_globals(m, specs)
    plans = [ProjectPlan(m, sp) for sp in specs]
    write_globals(m, specs)
    for pl in plans:
        write_project(m, pl)
    # harness settings
    g = m.ws("global")
    r = m.get("global", "dep_method").row
    g.cell(r, 5).value = dep_method
    m.wb.save(path)
    return m


def recalc(path):
    out = subprocess.run([sys.executable, RECALC, path, "120"], capture_output=True, text=True)
    try:
        j = json.loads(out.stdout.strip())
    except Exception:
        print(out.stdout, out.stderr)
        raise
    return j


def v29_values(rows, calc_rows=None):
    """rows: {line: row in Output_Standalone (IDR)}; calc_rows: {line: (sheet, row, sign)} alternative source
    (used when the standalone block is zeroed by a flag, e.g. TBM). Calculation sheets use I=2023, J..V=2024..2036."""
    wb = openpyxl.load_workbook(V29, data_only=True, read_only=True, keep_vba=True)
    res = {}
    if calc_rows:
        for name, spec in calc_rows.items():
            sheet, r, sign = spec
            ws = wb[sheet]
            res[name] = {y: sign * (ws.cell(r, c).value or 0) for y, c in V29_YEAR_COL.items()}
    ws = wb[STANDALONE]
    for name, r in rows.items():
        if r is None or name in res:
            continue
        res[name] = {y: (ws.cell(r, c).value or 0) for y, c in V29_YEAR_COL.items()}
    wb.close()
    return res


def compare(path, m, pid, tol):
    wb = openpyxl.load_workbook(path, data_only=True)
    sp = load_spec(pid)
    v29rows = sp.get("v29_rows", {})
    ref = v29_values(v29rows, sp.get("v29_calc_rows"))
    known = sp.get("v29_known_diffs", {})
    out = wb[m.title(f"{pid}:out")]
    keymap = {"rev": "revenue", "cogs": "cogs", "opex": "opex", "da": "da", "ebit": "ebit"}
    ok = True
    report = []
    for name, okey in keymap.items():
        if name not in ref:
            continue
        r = m.get(f"{pid}:out", okey).row
        line_ok = True
        diffs = []
        for y in range(2024, 2037):
            mine = out.cell(r, YEAR_COL[y]).value or 0
            theirs = ref[name][y] or 0
            d = (mine or 0) - (theirs or 0)
            if abs(d) > tol:
                line_ok = False
            diffs.append((y, mine, theirs, d))
        if not line_ok and name in known:
            report.append((name, "KNOWN", diffs))
            continue
        ok &= line_ok
        report.append((name, line_ok, diffs))
    # also EBIT vs v29 EBIT (v29 EBIT row = ebit key if given)
    print(f"\n=== {pid} {sp['name']} — comparison with v29 {STANDALONE} (tolerance {tol}) ===")
    for name, line_ok, diffs in report:
        tag = "MATCH" if line_ok is True else ("KNOWN DIFF (documented v29 source issue): " + known.get(name, "") if line_ok == "KNOWN" else "DIFF ")
        print(f"{name:6s} {tag}")
        if line_ok is not True:
            for y, mine, theirs, d in diffs:
                flag = "" if abs(d) <= tol else "  <-- "
                print(f"     {y}: model {mine:14.6f}  v29 {theirs:14.6f}  diff {d:12.6f}{flag}")
    # formula errors in the mini workbook
    errs = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#"):
                    errs.append(f"{ws.title}!{c.coordinate}={c.value}")
    if errs:
        ok = False
        print(f"FORMULA ERRORS ({len(errs)}): " + ", ".join(errs[:40]))
    return ok


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pid")
    ap.add_argument("--dep-method", type=int, default=1)
    ap.add_argument("--tol", type=float, default=0.001)
    ap.add_argument("--keep", action="store_true")
    a = ap.parse_args()
    pids = resolve_deps(a.pid)
    outdir = os.path.join(ROOT, "verify", "out")
    os.makedirs(outdir, exist_ok=True)
    path = os.path.join(outdir, f"verify_{a.pid}.xlsx")
    m = build(pids, path, dep_method=a.dep_method)
    j = recalc(path)
    print("recalc:", {k: j.get(k) for k in ("status", "total_formulas", "total_errors")})
    if j.get("error_summary"):
        print(json.dumps(j["error_summary"], indent=1)[:3000])
    ok = compare(path, m, a.pid, a.tol)
    print("RESULT:", "PASS" if ok else "FAIL")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
