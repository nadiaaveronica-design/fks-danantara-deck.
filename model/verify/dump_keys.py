#!/usr/bin/env python3
"""Print recalculated values of registered rows: python3 model/verify/dump_keys.py <xlsx> <sheetkey:key,...> [--only ids]
Rebuilds the registry (fast) with the same spec list used for the workbook, then reads cached values."""
import sys, os, argparse, warnings
warnings.filterwarnings("ignore")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import openpyxl
from fsl.engine import YEARS, YEAR_COL, L, SCALAR_COL
import build_model


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("keys", help="comma-separated sheetkey:key, e.g. fin:lt_close,fund:req_sel")
    ap.add_argument("--only", default=None)
    ap.add_argument("--years", default="2023-2036")
    a = ap.parse_args()
    only = a.only.split(",") if a.only else None
    tmp = a.xlsx + ".registry.tmp.xlsx"
    m = build_model.build(tmp, only=only)
    os.remove(tmp)
    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    y0, y1 = [int(x) for x in a.years.split("-")]
    for tok in a.keys.split(","):
        sk, key = tok.rsplit(":", 1)
        ref = m.get(sk, key)
        ws = wb[m.title(sk)]
        if ref.kind == "scalar":
            v = ws.cell(ref.row, ref.col).value
            print(f"{tok:32s} [{ref.label[:60]}] = {v}")
        else:
            vals = []
            for y in YEARS:
                if y0 <= y <= y1:
                    v = ws.cell(ref.row, YEAR_COL[y]).value
                    vals.append(f"{y}:{v:,.2f}" if isinstance(v, (int, float)) else f"{y}:{v}")
            print(f"{tok:32s} [{ref.label[:60]}]")
            print("      " + "  ".join(vals))


if __name__ == "__main__":
    main()
